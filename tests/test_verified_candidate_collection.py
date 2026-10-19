"""Raw-data collection must never turn partial or mixed-phase evidence into predictions."""
import importlib.util
from pathlib import Path
from types import SimpleNamespace
import pytest
ROOT=Path(__file__).resolve().parents[1]


def module():
 spec=importlib.util.spec_from_file_location('verified_collect_test',ROOT/'scripts/pipeline/collect_verified_candidates.py')
 m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);return m


def state_queue(tmp_path, statuses):
 import json
 tasks=[dict(key=f'm__s{i}',id='m',state=f's{i}') for i in range(len(statuses))]
 (tmp_path/'plan.json').write_text(json.dumps(dict(tasks=tasks)))
 for task,status in zip(tasks,statuses):
  if status is None:continue
  path=tmp_path/'tasks'/task['key'];path.mkdir(parents=True)
  (path/'status.json').write_text(json.dumps(dict(task=task,status=status)))
 return tasks


def test_native_state_queue_reports_pending_and_producer_progress_separately(tmp_path):
 state_queue(tmp_path,['state_records_verified','computing',None])
 progress,terminal=module().queue_progress(tmp_path)
 assert not terminal
 assert '1/3 surface records verified' in progress['m']
 assert '1 producer reports in progress' in progress['m']
 assert '1 queued/unresolved' in progress['m']
 assert 'field-specific QC separate' in progress['m']


def test_terminal_native_state_queue_does_not_hide_preserved_failures(tmp_path):
 state_queue(tmp_path,['state_records_verified','failed_preserved'])
 progress,terminal=module().queue_progress(tmp_path)
 assert terminal and '1 failed/preserved' in progress['m']
 assert '1/2 surface records verified' in progress['m']


def test_native_queue_never_counts_a_mismatched_state_report_as_complete(tmp_path):
 import json
 tasks=state_queue(tmp_path,['state_records_verified'])
 path=tmp_path/'tasks'/tasks[0]['key']/'status.json'
 path.write_text(json.dumps(dict(task=dict(key='other__s0',id='other'),status='state_records_verified')))
 progress,terminal=module().queue_progress(tmp_path)
 assert not terminal and '0/1 surface records verified' in progress['m']


def test_optional_state_QC_fields_survive_both_csv_and_workbook(tmp_path):
 import csv
 from openpyxl import load_workbook
 coverage=module().coverage
 rows=[dict(id='ordinary',G=-10.),dict(id='ensemble',G=-11.,single_surface_kinetics_validated=False)]
 csv_path=tmp_path/'states.csv';xlsx=tmp_path/'states.xlsx'
 coverage.write_csv(csv_path,rows)
 with csv_path.open() as stream:records=list(csv.DictReader(stream))
 assert records[0]['single_surface_kinetics_validated']==''
 assert records[1]['single_surface_kinetics_validated']=='False'
 assert coverage.workbook(xlsx,{'States':rows},{})==6
 sheet=load_workbook(xlsx)['States']
 assert sheet.cell(1,3).value=='single_surface_kinetics_validated'
 assert sheet.cell(2,3).value is None
 assert sheet.cell(3,3).value is False


@pytest.mark.parametrize('metadata', [None, {}, {'status':'hold_native_thermal_properties'}])
def test_thermal_method_review_withdraws_derived_properties_but_preserves_descriptors(tmp_path, metadata):
 import json
 m=module();path=tmp_path/'review.json'
 if metadata is not None:path.write_text(json.dumps(metadata))
 row=dict(E1_vs_Fc_V=-1.,E2_vs_Fc_V=-2.,E1_surface_estimate_V=-1.,
          E2_surface_estimate_V=-2.,wave_spacing_V=1.,dG_disproportionation_kJ_mol=96.,
          accessible_capacity_mAh_g=100.,capacity_nominal_mAh_g=150.,MW=350.,
          active_SP_states=3,dft_geometry_states=3)
 state=dict(native_solution_G_eV=-10.,native_surface_G_eV=-10.)
 couple=dict(E_native_vs_Fc_V=-1.,E_native_surface_estimate_vs_Fc_V=-1.,gas_Nelsen_sum_eV=.8)
 reference=dict(reference_native_Fc_absolute_V=4.37)
 assert m.apply_thermal_publication_hold([row],[state],[couple],reference,path)
 assert all(row[k] is None for k in ('E1_vs_Fc_V','E2_vs_Fc_V','wave_spacing_V',
            'dG_disproportionation_kJ_mol','accessible_capacity_mAh_g',
            'E1_surface_estimate_V','E2_surface_estimate_V'))
 assert row['MW']==350. and row['capacity_nominal_mAh_g']==150.
 assert row['active_SP_states']==row['dft_geometry_states']==3
 assert state['native_solution_G_eV'] is state['native_surface_G_eV'] is None
 assert couple['E_native_vs_Fc_V'] is couple['E_native_surface_estimate_vs_Fc_V'] is None
 assert couple['gas_Nelsen_sum_eV']==.8 and reference['reference_native_Fc_absolute_V'] is None


@pytest.mark.parametrize('failure',['state','missing_spacing'])
def test_independent_native_disagreement_quarantines_derived_properties(monkeypatch,failure):
 m=module()
 values={('new','neu'):-9.,('new','red1'):-14.,('new','red2'):-18.,('ferrocene','ox'):-4.,('ferrocene','neu'):-10.}
 monkeypatch.setattr(m,'independent_state',lambda gid,st,q,key,audit:(values[gid,st],'admitted',{}))
 monkeypatch.setattr(m,'read_manifest',lambda:[dict(id='ferrocene',state='ox',smiles='CC'),dict(id='ferrocene',state='neu',smiles='CC')])
 rows=[dict(id='p',registry_id='new',exact_grafted_smiles='CCO',n_e_nominal=2,capacity_nominal_mAh_g=1163.6,
      states='[["neu",0,1,0],["red1",-1,2,-1],["red2",-2,1,-2]]',E1_vs_Fc_V=-1,E2_vs_Fc_V=-2,
      wave_spacing_V=1 if failure=='state' else None,dG_disproportionation_kJ_mol=None,accessible_capacity_mAh_g=None,
      solution_potentials_status='verified native calculation',result_stage='native solution states complete',
      native_solution_ready_states=3,accessible_capacity_status='pending')]
 states=[dict(proposal_id='p',registry_id='new',state=s,requested_charge=q,exact_grafted_inchikey='key',
       native_solution_G_eV=g,native_solution_status='verified calculation') for s,q,g in [('neu',0,-9),('red1',-1,-13.9 if failure=='state' else -14),('red2',-2,-18)]]
 couples=[dict(proposal_id='p',registry_id='new',event='neu->red1',E_native_vs_Fc_V=-1,
      gas_Nelsen_status='flagged diagnostic',gas_Nelsen_sum_eV=.9)]
 audit=SimpleNamespace(FARADAY=96485.33212)
 before=set(rows[0])
 checks,_,_=m.raw_audit_and_quarantine(rows,states,couples,audit)
 assert any(not c['passed'] for c in checks)
 assert rows[0]['E1_vs_Fc_V'] is None and rows[0]['E2_vs_Fc_V'] is None
 if failure=='state':assert states[1]['native_solution_G_eV'] is None
 else:assert any(c['field']=='wave_spacing_V' and not c['passed'] and c['independent_value']==1 for c in checks)
 assert couples[0]['gas_Nelsen_sum_eV'] is None  # flagged gas diagnostic never becomes a solution prediction
 assert set(rows[0])==before  # all rows remain exportable under the same CSV schema


def test_queue_done_marker_is_progress_not_an_independent_metric(tmp_path):
 import json
 m=module();(tmp_path/'queue.json').write_text(json.dumps(dict(ids=['a','b'])))
 (tmp_path/'done').mkdir();(tmp_path/'done/a.json').write_text('{}')
 progress,terminal=m.queue_progress(tmp_path)
 assert not terminal and 'record QC separate' in progress['a'] and progress['b']=='pending'
 (tmp_path/'failures').mkdir();(tmp_path/'failures/b.json').write_text(json.dumps(dict(attempts=2)))
 progress,terminal=m.queue_progress(tmp_path)
 assert terminal and progress['b'].startswith('terminal failure')


def test_unresolved_small_spin_gap_keeps_clean_first_wave_but_blocks_full_molecule(monkeypatch):
 m=module()
 def spin(gid,st):
  return dict(spin_confidence='DFT_check_required' if st=='red2' else 'UMA_large_gap_preliminary',
     spin_gap_eV=.2 if st=='red2' else 2.,spin_gap_source='UMA',spin_evidence_path=None,
     spin_primary_admissible=st!='red2')
 monkeypatch.setattr(m,'state_spin_evidence',spin)
 states=[dict(proposal_id='p',registry_id='p',state=st,native_solution_G_eV=g,
               native_solution_status='verified calculation') for st,g in [('neu',-9),('red1',-14),('red2',-18)]]
 couples=[dict(proposal_id='p',event=ev,q_ox=q,E_native_vs_Fc_V=e,potential_status='native calculated/QC passed')
          for ev,q,e in [('neu->red1',0,-1),('red1->red2',-1,-2)]]
 rows=[dict(id='p',declared_states=3,states='[["neu",0,1,0],["red1",-1,2,-1],["red2",-2,1,-2]]',
           E1_vs_Fc_V=-1,E2_vs_Fc_V=-2,dG_disproportionation_kJ_mol=96.485,wave_spacing_V=1,
           accessible_capacity_mAh_g=100,result_stage='native solution states complete',
           disproportionation_status='native thermodynamic calculation')]
 m.apply_spin_gates(rows,states,couples)
 assert rows[0]['E1_vs_Fc_V']==-1 and rows[0]['E2_vs_Fc_V'] is None
 assert rows[0]['spin_confidence']=='DFT_check_required' and rows[0]['accessible_capacity_mAh_g'] is None
 assert rows[0]['dG_disproportionation_kJ_mol'] is None
 assert states[-1]['native_surface_G_eV']==-18 and states[-1]['native_solution_G_eV'] is None
 assert rows[0]['native_surface_calculation_states']==3 and rows[0]['native_solution_ready_states']==2


def test_ready_empty_issue_field_roundtrips_as_an_excel_blank(tmp_path):
 m=module()
 count=m.coverage.workbook(tmp_path/'snapshot.xlsx',{'States':[dict(id='p',state_issues='',native_solution_G_eV=None)]},{})
 assert count==3


@pytest.mark.parametrize('oxidation,inverted,aliases',[(False,False,False),(True,False,False),(True,True,False),(True,False,True)])
def test_wave_spacing_is_signed_for_both_ladder_directions_without_reference(monkeypatch,oxidation,inverted,aliases):
 import json
 m=module()
 monkeypatch.setattr(m,'state_spin_evidence',lambda *args:dict(spin_confidence='admissible',spin_gap_eV=1.,
      spin_gap_source='test',spin_evidence_path=None,spin_primary_admissible=True))
 names=['neu','ox1','ox2'] if oxidation else ['neu','red1','red2'];charges=[0,1,2] if oxidation else [0,-1,-2]
 energies=[-100.,-95.,-91. if inverted else -89.] if oxidation else [-100.,-105.,-109.]
 actual=['actual0','actual1','actual2'] if aliases else names
 states=[dict(proposal_id='p',registry_id='p',state=st,requested_charge=q,native_solution_G_eV=g,native_solution_status='verified calculation') for st,q,g in zip(actual,charges,energies)]
 couples=[]
 for i in range(2):
  event=actual[i+1]+'->'+actual[i] if oxidation else actual[i]+'->'+actual[i+1]
  couples.append(dict(proposal_id='p',event=event,q_ox=charges[i+1] if oxidation else charges[i],E_native_vs_Fc_V=None))
 definition=[(st,q,2 if abs(q)==1 else 1,q) for st,q in zip(names,charges)]
 row=dict(id='p',declared_states=3,states=json.dumps(definition),dG_disproportionation_kJ_mol=None)
 m.apply_spin_gates([row],states,couples)
 assert row['E1_vs_Fc_V'] is None and row['E2_vs_Fc_V'] is None
 assert row['wave_spacing_V']==pytest.approx(-1. if inverted else 1.)


@pytest.mark.parametrize('oxidation',[False,True])
def test_missing_first_wave_never_renumbers_second_wave_surface(monkeypatch,oxidation):
 m=module()
 monkeypatch.setattr(m,'state_spin_evidence',lambda gid,st:dict(spin_confidence='standard_model',spin_primary_admissible=True))
 direction=1 if oxidation else -1
 names=['neu','ox1','ox2'] if oxidation else ['neu','red1','red2']
 declared=[[st,i*direction,1,i*direction] for i,st in enumerate(names)]
 states=[dict(proposal_id='p',registry_id='p',state=st,native_solution_G_eV=None if i==0 else -i,
              native_solution_status='pending' if i==0 else 'verified calculation') for i,st in enumerate(names)]
 events=[names[i+1]+'->'+names[i] if oxidation else names[i]+'->'+names[i+1] for i in range(2)]
 couples=[dict(proposal_id='p',event=ev,q_ox=(i+1 if oxidation else -i),E_native_vs_Fc_V=None if i==0 else .7,
               potential_status='pending' if i==0 else 'native calculated/QC passed') for i,ev in enumerate(events)]
 import json
 rows=[dict(id='p',states=json.dumps(declared),declared_states=3,E1_vs_Fc_V=.7,E2_vs_Fc_V=None,
            result_stage='partial',solution_potentials_status='pending')]
 m.apply_spin_gates(rows,states,couples)
 assert rows[0]['E1_vs_Fc_V'] is None and rows[0]['E1_surface_estimate_V'] is None
 assert rows[0]['E2_vs_Fc_V']==.7 and rows[0]['E2_surface_estimate_V']==.7


def test_near_DFT_spin_gaps_require_native_thermals_or_spin_ensemble():
 m=module()
 assert m.qualify_DFT_spin_gap(.15,False)==('matching_native_spin_thermals_required',False)
 assert m.qualify_DFT_spin_gap(-.08,False)==('matching_native_spin_thermals_required',False)
 assert m.qualify_DFT_spin_gap(.08,True)==('spin_ensemble_required',False)
 assert m.qualify_DFT_spin_gap(-.08,True)==('spin_ensemble_required',False)
 assert m.qualify_DFT_spin_gap(.7,False)==('DFT_electronic_checked_thermal_uncertainty',True)
 assert m.qualify_DFT_spin_gap(-.7,False)==('wrong_ground_state',False)


@pytest.mark.parametrize('changed_gid',['a','ferrocene',None])
def test_inflight_input_quarantines_its_dependents_and_keeps_other_results(monkeypatch,changed_gid):
 m=module();values={('ferrocene','ox'):-4,('ferrocene','neu'):-10}
 for gid in ['a','b']:
  values.update({(gid,st):g for st,g in [('neu',-9),('red1',-14),('red2',-18)]})
 monkeypatch.setattr(m,'independent_state',lambda gid,st,*args:(values[gid,st],'admitted',{}))
 monkeypatch.setattr(m,'read_manifest',lambda:[dict(id='ferrocene',state='ox',smiles='CC'),dict(id='ferrocene',state='neu',smiles='CC')])
 rows=[dict(id=gid,registry_id=gid,exact_grafted_smiles='CCO',n_e_nominal=2,capacity_nominal_mAh_g=1163.5,
    states='[["neu",0,1,0],["red1",-1,2,-1],["red2",-2,1,-2]]',E1_vs_Fc_V=-1,E2_vs_Fc_V=-2,
    wave_spacing_V=1,dG_disproportionation_kJ_mol=None,accessible_capacity_mAh_g=None,
    solution_potentials_status='verified native calculation',result_stage='native solution states complete',
    native_solution_ready_states=3,accessible_capacity_status='pending') for gid in ['a','b']]
 states=[dict(proposal_id=gid,registry_id=gid,state=st,requested_charge=q,exact_grafted_inchikey='key',
    native_solution_G_eV=g,native_solution_status='verified calculation')
    for gid in ['a','b'] for st,q,g in [('neu',0,-9),('red1',-1,-14),('red2',-2,-18)]]
 changed=m.ROOT/('config/new.py' if changed_gid is None else f'calcs/spincheck/{changed_gid}/red2_m3/result.json')
 context={};checks,_,changing=m.raw_audit_and_quarantine(rows,states,[],SimpleNamespace(FARADAY=96485.33212),
    reference=context,input_inventory=(set(),lambda:{changed}))
 assert all(c['passed'] for c in checks)
 assert changing==[str(changed)]
 assert rows[0]['E1_vs_Fc_V'] is None
 if changed_gid=='a':
  assert rows[1]['E1_vs_Fc_V']==-1 and rows[1]['native_solution_ready_states']==3
  assert context['quarantined_input_gids']==['a'] and not context['global_input_changed']
 else:
  assert rows[1]['E1_vs_Fc_V'] is None and context['global_input_changed']


def test_unfinished_positive_spin_helper_never_admits_primary_gap():
 m=module();identity={'geometry_sha256':'current'}
 report=dict(status='spin_gap_verified',chosen_identity=identity,electronic_gap_eV=.8)
 for status in ['pending','failed','source_stale']:
  assert m.helper_spin_report_admission(report,{'spin_validation_status':status},identity)[1] is False
 for status in ['running','near_gap_waiting_chosen_native','lower_spin_waiting_motif_release','failed','source_stale']:
  assert m.helper_spin_report_admission(dict(report,status=status),{'spin_validation_status':'validated'},identity)[1] is False
 assert m.helper_spin_report_admission(report,{'spin_validation_status':'validated'},identity)[1] is True
 assert m.helper_spin_report_admission(report,{'spin_validation_status':'validated'},{'geometry_sha256':'changed'})[1] is False
 assert m.helper_spin_report_admission(dict(report,ensemble_primary_integration_pending=True),{'spin_validation_status':'validated'},identity)[1] is False
