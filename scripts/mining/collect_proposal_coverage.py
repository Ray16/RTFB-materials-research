#!/usr/bin/env python
"""Collect exact-structure results for the original39 proposals without parent inheritance.

Writes a fresh CSV/XLSX snapshot. All property columns come from addressed grafted
calculations; parent/database quantities remain separate provenance, never substitutes.
"""
from __future__ import annotations
import argparse,csv,hashlib,json,math,sys,time
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'src'))
from rdkit import Chem
from rdkit.Chem import Descriptors,rdMolDescriptors
from redox.build.contracts import validate_group
from redox.core.common import load_config,read_manifest,read_result,UMA,DFT,free_energy
from redox.screening.discovery import describe,precursor_from_frag
from redox.qm.uma import geometry_checks
from redox.qm.native_thermal import load_native_thermal
from redox.validation.readiness import inspect_states
from redox.properties import integrity,reorg
from redox.screening.scorecard import contiguous_paths
CONFIGS=('redox_groups','starting_candidates','merrifield_multielectron','discovered_candidates','custom_candidates')


def safe_json(path):
 try:return json.loads(Path(path).read_text())
 except(FileNotFoundError,ValueError,OSError):return None


def full_key(smi):
 return Chem.MolToInchiKey(Chem.MolFromSmiles(smi))


def state_rows(group,key,manifest,registered):
 rows=[];gids={r['id'] for r in manifest if full_key(r['smiles'])==key}
 if len(gids)>1:raise ValueError('multiple library aliases for exact proposal graph: '+group['id'])
 gid=next(iter(gids),None)
 matched=[r for r in manifest if r['id']==gid] if gid else []
 for label,q,m,ne in group['states']:
  actual=next((r for r in matched if int(r['charge'])==q and int(r['n_e'])==ne),None)
  state=actual['state'] if actual else label
  uma=safe_json(UMA/gid/state/'result.json') if gid and actual else None
  dft=safe_json(DFT/gid/state/'result.json') if gid and actual else None
  graph=False
  if uma and uma.get('smiles') and full_key(uma['smiles'])==key:
   try:graph=not geometry_checks(uma['smiles'],UMA/gid/state/'relaxed.xyz',q,uma['chosen_mult'])
   except(Exception):graph=False
  uok=bool(uma and graph and uma.get('charge')==q and uma.get('n_e')==ne and uma.get('converged') is True
       and isinstance(uma.get('energy_eV'),(int,float)) and math.isfinite(uma['energy_eV']))
  read=read_result(gid,state) if gid and actual and dft else None
  native=load_native_thermal(DFT/gid/state,q,dft['mult']) if dft else None
  readiness=inspect_states([actual])[0] if actual else dict(status='NOT_BUILT',issues='no exact registered graph/state')
  G,tq=free_energy(read)
  # Admit solution predictions only under native/current thermal and geometry QC.
  solok=bool(read and readiness['status']=='ready' and read.get('thermal_definition')=='native_solution_dft_qrrho_spin'
             and tq=='ok' and G is not None and read.get('sp_status')=='ok')
  rows.append(dict(proposal_id=group['id'],exact_grafted_inchikey=key,registry_id=gid,state=state,
     requested_charge=q,requested_mult_hint=m,n_e=ne,manifest_state_present=actual is not None,
     uma_gas_geometry_status='verified graph/warm start' if uok else ('invalid/unverified' if uma else 'pending'),
     uma_cache_schema=uma.get('cache_schema') if uma else None,
     uma_cache_provenance='fingerprinted' if uma and uma.get('model_identity') and uma.get('seed_identity') and uma.get('geometry_sha256') else ('legacy; rerun required' if uma else 'pending'),
     uma_gas_energy_eV=uma['energy_eV'] if uok else None,uma_mult=uma.get('chosen_mult') if uok else None,
     uma_model=uma.get('model') if uok else None,uma_fmax_eV_A=uma.get('fmax') if uok else None,
     dft_geometry_present=bool(dft and (DFT/gid/state/'opt.xyz').exists()),
     active_sp_status=read.get('sp_status') if read else 'pending',
     active_smd_energy_eV=read.get('e_smd_eV') if read and read.get('sp_status')=='ok' else None,
     native_thermal_status='addressed QC passed' if native else 'pending',
     native_g_thermal_eV=native.get('g_thermal_total_eV') if native else None,
     effective_g_thermal_eV=read.get('g_thermal_eV') if read and tq=='ok' and read.get('thermal_definition')=='native_solution_dft_qrrho_spin' else None,
     spin_ensemble_status=read.get('spin_ensemble_status') if read else None,
     spin_ensemble_correction_eV=read.get('spin_ensemble_correction_eV') if read else None,
     spin_ensemble_alternative_population=read.get('spin_ensemble_alternative_population') if read else None,
     selected_thermal_definition=read.get('thermal_definition') if read else None,
     state_readiness=readiness['status'],state_issues=readiness['issues'],
     native_solution_G_eV=G if solok else None,native_solution_status='verified calculation' if solok else 'pending',
     result_path=str(DFT/gid/state/'result.json') if dft else None,
     uma_result_path=str(UMA/gid/state/'result.json') if uma else None))
 return rows,gid


def collect(group,manifest,registered,sources,faraday):
 mol=validate_group('Cc1ccc(C[*:1])cc1',group);smi=Chem.MolToSmiles(mol);key=full_key(smi)
 stats=describe(smi,max(abs(s[3]) for s in group['states']),faraday)
 atomic=sum(a.GetMass() for a in Chem.AddHs(mol).GetAtoms())
 assert abs(stats['MW']-atomic)<.00051
 states,gid=state_rows(group,key,manifest,registered)
 primary=registered.get(key)
 confs=list((ROOT/'library'/gid/'conformers').glob('conf_*.xyz')) if gid else []
 ns=len(states);warm=sum(s['uma_gas_geometry_status']=='verified graph/warm start' for s in states)
 native=sum(s['native_solution_status']=='verified calculation' for s in states)
 prow=group.get('provenance',{})
 urls='; '.join(sources[k]['url'] for k in prow.get('source_keys',[]))
 Fc=[]
 for st in ('ox','neu'):
  r=read_result('ferrocene',st);G,t=free_energy(r)
  Fc.append(G if r and r.get('thermal_definition')=='native_solution_dft_qrrho_spin' and t=='ok' else None)
 fcref=Fc[0]-Fc[1] if all(v is not None for v in Fc) else None
 ordered=sorted(states,key=lambda s:-s['requested_charge']);events=[]
 for ox,red in zip(ordered,ordered[1:]):
  Gox,Gred=ox['native_solution_G_eV'],red['native_solution_G_eV']
  if ox['requested_charge']-red['requested_charge']==1 and Gox is not None and Gred is not None and fcref is not None:
   events.append(dict(event=ox['state']+'->'+red['state'],q_ox=ox['requested_charge'],q_red=red['requested_charge'],
     E_vs_Fc_V=-(Gred-Gox)-fcref,reference='matched native Fc',definition='native solution free energy'))
 states_nominal=group['states'];rest=next(q for st,q,m,ne in states_nominal if ne==0)
 reduction=[e for e in events if e['q_ox']<=rest];oxidation=[e for e in events if e['q_red']>=rest]
 kind='oxidation' if max(s[3] for s in states_nominal)>0 and min(s[3] for s in states_nominal)==0 else 'reduction'
 # Preserve physical step positions even when one native wave is unavailable.
 pairs=[(ox,red) for ox,red in zip(ordered,ordered[1:])
        if (red['requested_charge']>=rest if kind=='oxidation' else ox['requested_charge']<=rest)]
 if kind=='oxidation':pairs=pairs[::-1]
 chosen=[next((e for e in events if e['event']==ox['state']+'->'+red['state']),
              dict(E_vs_Fc_V=None)) for ox,red in pairs]
 couples=[];assessed={}
 if gid and any(s['dft_geometry_present'] for s in states):
  try:assessed={r['couple']:r for r in integrity.assess(gid)}
  except (OSError,ValueError,KeyError,TypeError):pass
 event_map={e['event']:e for e in events}
 for ox,red in zip(ordered,ordered[1:]):
  name=ox['state']+'->'+red['state'];event=event_map.get(name);assessment=assessed.get(name)
  gas=None
  if gid and ox['dft_geometry_present'] and red['dft_geometry_present']:
   try:
    ro=read_result(gid,ox['state']);rr=read_result(gid,red['state'])
    if ro and rr:gas=reorg.lambda_for_couple(gid,(ox['state'],ox['requested_charge'],ro['mult']),
                                            (red['state'],red['requested_charge'],rr['mult']))
   except (OSError,ValueError,KeyError,TypeError):pass
  couples.append(dict(proposal_id=group['id'],registry_id=gid,exact_grafted_inchikey=key,event=name,
      q_ox=ox['requested_charge'],q_red=red['requested_charge'],
      E_native_vs_Fc_V=event['E_vs_Fc_V'] if event else None,
      potential_status='native calculated/QC passed' if event else 'pending',
      integrity_status=assessment['verdict'] if assessment else 'pending',
      solvated_EA_eV=assessment['EA_solv_eV'] if assessment else None,
      ensemble_integrity_status=assessment.get('ensemble_integrity_qc') if assessment else None,
      ensemble_component_pairs=assessment.get('ensemble_component_pairs') if assessment else None,
      gas_Nelsen_sum_eV=gas['lambda_i_eV'] if gas else None,
      gas_Nelsen_flags=gas['flag'] if gas else None,
      gas_Nelsen_status=('flagged diagnostic' if gas['flag'] else 'QC-clean gas diagnostic') if gas else 'pending',
      solution_relaxation_status='pending separate validated solution method',
      outer_sphere_status='pending separate validated solution method',parent_metrics_inherited=False))
 disp=None
 if len(ordered)==3 and all(s['native_solution_G_eV'] is not None for s in ordered):
  disp=(ordered[0]['native_solution_G_eV']+ordered[2]['native_solution_G_eV']-2*ordered[1]['native_solution_G_eV'])*faraday/1000
 accessible=None
 if len(events)==ns-1 and len(assessed)==ns-1:
  ele=load_config('electrolyte');lo,hi=ele.WINDOW_V_VS_FC
  path_input={c['q_ox']:dict(couple=c['event'],E=c['E_native_vs_Fc_V'],
      status='ok' if c['E_native_vs_Fc_V'] is not None else 'INCOMPLETE',integrity=c['integrity_status'],
      in_window=c['E_native_vs_Fc_V'] is not None and lo<=c['E_native_vs_Fc_V']<=hi) for c in couples}
  paths,_=contiguous_paths(path_input,rest,ele.ANOLYTE_CATHOLYTE_DIVIDER_V)
  pool='catholyte' if kind=='oxidation' else 'anolyte'
  accessible=len(paths[pool])*faraday/(3.6*stats['MW'])
 geometry_count=sum(s['dft_geometry_present'] for s in states)
 sp_count=sum(s['active_sp_status']=='ok' for s in states)
 thermal_count=sum(s['native_thermal_status']=='addressed QC passed' for s in states)
 if native==ns:stage='native solution states complete'
 elif native:stage='partial native solution states; remaining states pending'
 elif thermal_count:stage='native thermal records present; solution admission pending'
 elif sp_count:stage='active SP records present; native validation pending'
 elif geometry_count:stage='DFT geometry records present; native validation pending'
 elif warm:stage='gas UMA only'
 else:stage='built; computations pending' if confs else 'proposal; computations pending'
 row=dict(id=group['id'],name=group['name'],family=group['family'],origin=prow.get('origin'),
    review_priority=group.get('review_priority'),result_stage=stage,registry_id=gid,registered=primary is not None,
    exact_grafted_smiles=smi,exact_grafted_inchikey=key,formula=rdMolDescriptors.CalcMolFormula(mol),
    frag=group['frag'],precursor_smiles=precursor_from_frag(group['frag']),states=json.dumps(states_nominal),
    descriptors_status='verified exact-grafted graph/arithmetic',**stats,
    total_atoms_including_H=Chem.AddHs(mol).GetNumAtoms(),rotatable_bonds=Descriptors.NumRotatableBonds(mol),
    declared_states=ns,canonical_conformer_count=len(confs),uma_gas_verified_states=warm,
    dft_geometry_states=geometry_count,
    active_SP_states=sp_count,native_thermal_QC_states=thermal_count,
    native_solution_ready_states=native,solution_potentials_status='verified native calculation' if len(chosen)==ns-1 and all(e['E_vs_Fc_V'] is not None for e in chosen) else 'pending',
    E1_vs_Fc_V=chosen[0]['E_vs_Fc_V'] if chosen else None,E2_vs_Fc_V=chosen[1]['E_vs_Fc_V'] if len(chosen)>1 else None,
    wave_spacing_V=(chosen[0]['E_vs_Fc_V']-chosen[1]['E_vs_Fc_V'])*(1 if kind=='reduction' else -1) if len(chosen)>1 and all(e['E_vs_Fc_V'] is not None for e in chosen) else None,
    dG_disproportionation_kJ_mol=disp,
    disproportionation_status='not applicable one-electron ladder' if ns<3 else ('native thermodynamic calculation; decomposition untested' if disp is not None else 'pending'),
    reorganization_status='gas diagnostic available; solution descriptor pending' if any(c['gas_Nelsen_sum_eV'] is not None for c in couples) else 'pending',
    lambda_i_eV=None,lambda_outer_eV=None,
    accessible_capacity_status='computed native-window/integrity bookkeeping; kinetic QC separate' if accessible is not None else 'pending electrolyte-window/integrity QC',
    accessible_capacity_mAh_g=accessible,
    experimental_grafted_properties_verified=False,parent_metrics_inherited=False,
    source_urls=urls or (primary.get('provenance',{}).get('source') if primary else None),
    chemistry_notes=group.get('chemistry_notes'),remaining_work='; '.join(
       issue for condition,issue in [(not primary,'register exact graph/states'),(not confs,'build conformer ensemble'),
         (warm<ns,'UMA complete cache/graph evidence'),(geometry_count<ns,'remaining DFT geometry records'),
         (sp_count<ns,'remaining active SP records'),(native<ns,'native Hessian/derivative and solution-admission QC'),
         (len(chosen)<ns-1 or any(e['E_vs_Fc_V'] is None for e in chosen),'matched-reference redox potentials'),(True,'field-specific integrity/reorganization and window gates')] if condition))
 return row,states,couples


def columns(rows):
 """Preserve optional QC fields wherever they first appear in the snapshot."""
 return list(dict.fromkeys(key for row in rows for key in row))


def write_csv(path,rows):
 from redox.core.io import atomic_text_writer
 fields=columns(rows)
 with atomic_text_writer(path,newline='') as f:
  if fields:
   w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)


def workbook(path,sheets,notes):
 from openpyxl import Workbook,load_workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 from openpyxl.utils import get_column_letter
 from openpyxl.comments import Comment
 wb=Workbook();wb.remove(wb.active)
 for name,rows in sheets.items():
  fields=columns(rows);ws=wb.create_sheet(name);ws.append(fields)
  for r in rows:ws.append([r.get(k) for k in fields])
  ws.freeze_panes='E2';ws.auto_filter.ref=ws.dimensions
  ws.row_dimensions[1].height=42
  for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173F5F');cell.alignment=Alignment(wrap_text=True)
  for i,key in enumerate(fields,1):
   ws.column_dimensions[get_column_letter(i)].width=35 if 'status' in key or 'stage' in key else (55 if any(x in key for x in ('smiles','source','notes','work')) else 22)
   ws.cell(1,i).comment=Comment(notes.get(key,'Snapshot exported from exact-structure pipeline records; blank means not computed/validated.'),'Source/QC')
  for row in ws.iter_rows(min_row=2):
   for cell in row:
    cell.alignment=Alignment(vertical='top',wrap_text=True)
    if isinstance(cell.value,float):cell.number_format='0.0000'
    if cell.value in ('pending','proposal; computations pending','legacy; rerun required'):cell.fill=PatternFill('solid',fgColor='FFF2CC')
 ws=wb.create_sheet('Conventions')
 for key,value in notes.items():ws.append([key,value])
 ws.column_dimensions['A'].width=35;ws.column_dimensions['B'].width=100
 wb.save(path)
 read=load_workbook(path,data_only=False)
 count=0
 for name,rows in sheets.items():
  ws=read[name];assert ws.max_row==len(rows)+1
  for i,r in enumerate(rows,2):
   for j,key in enumerate(columns(rows),1):
    want=r.get(key);have=ws.cell(i,j).value
    assert have==want or (want=='' and have is None) or (isinstance(want,float) and isinstance(have,(float,int)) and abs(have-want)<1e-10),(name,i,key,have,want)
    assert ws.cell(i,j).data_type!='e';count+=1
 return count


def main():
 ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--out',required=True,type=Path);a=ap.parse_args()
 if a.out.exists():ap.error('choose a fresh snapshot directory')
 base=ROOT/'results/discovery/parallel_expansion_2026-10-10'
 proposals=json.loads((base/'proposals.json').read_text());assert len(proposals)==39
 sources=json.loads((base/'sources.json').read_text());manifest=read_manifest();registered={}
 for config in CONFIGS:
  cfg=load_config(config)
  for g in cfg.GROUPS:
   m=validate_group(cfg.SCAFFOLD,g);registered[Chem.MolToInchiKey(m)]=g
 faraday=load_config('electrolyte').FARADAY
 rows=[];states=[];couples=[]
 for group in proposals:
  row,detail,couple=collect(group,manifest,registered,sources,faraday);rows.append(row);states+=detail;couples+=couple
 extras=[];original_keys={r['exact_grafted_inchikey'] for r in rows}
 for g in load_config('custom_candidates').GROUPS:
  mol=validate_group('Cc1ccc(C[*:1])cc1',g)
  if Chem.MolToInchiKey(mol) not in original_keys:
   row,detail,couple=collect(g,manifest,registered,sources,faraday);extras.append(row);states+=detail;couples+=couple
 a.out.mkdir(parents=True)
 write_csv(a.out/'original_39_coverage.csv',rows);write_csv(a.out/'state_coverage.csv',states);write_csv(a.out/'couple_coverage.csv',couples)
 if extras:write_csv(a.out/'extra_registered_candidates.csv',extras)
 notes={'snapshot_utc':datetime.now(ZoneInfo('UTC')).isoformat(),'snapshot_Chicago':datetime.now(ZoneInfo('America/Chicago')).isoformat(),
  'MW':'Exact grafted active-species mass (g/mol). No counterions or polymer repeat correction.',
  'capacity_nominal_mAh_g':'Declared electrons * F/(3.6*MW); graph bookkeeping, not computed usable-electron capacity.',
  'SA_grafted':'RDKit Ertl heuristic of the exact graft; not reaction selectivity or measured synthetic cost.',
  'solution_potentials_status':'Admitted only from exact graph/states, active SMD records, native thermal QC and matched native Fc.',
  'uma_gas_verified_states':'Geometry/charge/spin parity verified; gas warm starts are not solution properties. Legacy cache must rerun.',
  'parent_metrics_inherited':'Always false. No D3TaLES parent or literature value becomes a grafted prediction.',
  'empty_numeric_cells':'Metric not computed or not admissible; never a substituted zero.',
  'workbook':'Immutable results snapshot imported from audited CSV data; no spreadsheet formulas or cached formula errors.'}
 sheets={'Original39':rows,'StateCoverage':states,'CoupleCoverage':couples}
 if extras:sheets['ExtraRegistered']=extras
 audited=workbook(a.out/'original_39_results_coverage.xlsx',sheets,notes)
 summary=dict(**notes,original_proposals=len(rows),original_unique_connectivities=len({r['exact_grafted_inchikey'][:14] for r in rows}),
   original_registered=sum(r['registered'] for r in rows),original_built=sum(r['canonical_conformer_count']>0 for r in rows),
   original_gas_UMA_molecules=sum(r['uma_gas_verified_states']>0 for r in rows),original_gas_UMA_states=sum(r['uma_gas_verified_states'] for r in rows),
   original_DFT_molecules=sum(r['dft_geometry_states']>0 for r in rows),original_native_solution_complete_molecules=sum(r['native_solution_ready_states']==r['declared_states'] for r in rows),
   original_native_potential_molecules=sum(r['solution_potentials_status']=='verified native calculation' for r in rows),
   extra_registered_ids=[r['id'] for r in extras],workbook_cells_verified=audited,
   source_sha256={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [base/'proposals.json',ROOT/'config/custom_candidates.py',ROOT/'library/manifest.csv']})
 (a.out/'coverage_summary.json').write_text(json.dumps(summary,indent=2)+'\n');print(json.dumps(summary,indent=2))
if __name__=='__main__':main()
