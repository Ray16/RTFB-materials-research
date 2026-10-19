"""Independent read-only audit: re-derive every computable screening axis from the raw energy
records and compare it with the published tables in results/.

INDEPENDENCE: this script does NOT import redox.properties.*, redox.screening.* or
redox.validation.* (the code under audit). It uses only the energy-data contract
(redox.core.protocol: record_path / load_record / protocol_hash / geom_sha1 / ACTIVE_SP),
config/*.py, library/manifest.csv and the files under calcs/dft/. All arithmetic (free
energies, E°, ΔG_disp, Nelsen λ, capacity, contiguous paths, σ derivation, Pareto dominance)
is re-implemented here from the documented definitions (docs/DESIGN_AXES.md + module
docstrings).

Outputs (the only files written):
  results/validation/audit_recompute.csv  one row per checked quantity
  results/validation/audit_sanity.csv     physical plausibility flags on the raw records

Verdicts in audit_recompute.csv:
  ok           |published - recomputed| <= tol (tol = 1e-6 on the value rounded to the
               published precision, i.e. an exact recompute of the published digits)
  DISCREPANCY  anything else (incl. a value present on one side only)
check_type: exact (re-derived from raw records) | trace (number copied from another table,
            checked against that table / against our recompute) | logic (status, membership,
            flags) | cross_check (consistency between two published tables; tol stated)

Sanity flags in audit_sanity.csv are JUDGEMENT, not exact checks: each row says whether its
threshold is `heuristic` or `computed` (e.g. a Born estimate from the geometry).

  source ~/miniforge3/etc/profile.d/conda.sh && conda activate redox
  OMP_NUM_THREADS=1 PYTHONPATH=src python scripts/validation/audit/recompute_axes.py
"""
from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import math
import os
import statistics
import sys
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "src"))
from redox.core.protocol import ACTIVE_SP, geom_sha1, load_record, protocol_hash  # noqa: E402

DFT = ROOT / "calcs" / "dft"
RES = ROOT / "results"
OUT = RES / "validation"

# physical constants (own copies — not imported from the audited code)
EV_KJMOL = 96.485                     # documented eV -> kJ/mol used for the tables
KT_EV_298 = 0.0256926                 # k_B T at 298.15 K (eV)
COULOMB_EV_A = 14.399645              # e^2 / (4 pi eps0) in eV*A
F_TABLE_CAPACITY = 96485.0            # F used by capacity_and_proxies (docstring: 96485 C/mol)
C_PER_MAH = 3.6
C_MASS = 12.011

TOL = 1e-6


# ----------------------------------------------------------------------------- config / io
def cfg(name):
    spec = importlib.util.spec_from_file_location(name, ROOT / "config" / f"{name}.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def read_csv(p):
    p = Path(p)
    if not p.exists():
        return []
    with p.open() as f:
        return list(csv.DictReader(f))


def fnum(x):
    try:
        if x is None or (isinstance(x, str) and x.strip() == ""):
            return None
        v = float(x)
        return v if math.isfinite(v) else None
    except (TypeError, ValueError):
        return None


def rnd(x, d):
    return None if x is None else round(x, d)


ELE = cfg("electrolyte")
SCC = cfg("scorecard_config")
FARADAY = ELE.FARADAY
MANIFEST = read_csv(ROOT / "library" / "manifest.csv")


def is_state_dir(sd: Path):
    return (sd.is_dir() and not sd.name.startswith(("_", ".")) and ".bak" not in sd.name
            and sd.name not in ("reorg", "sp")
            and (sd / "result.json").exists() and (sd / "opt.xyz").exists())


def calc_ids():
    return sorted(d.name for d in DFT.iterdir()
                  if d.is_dir() and ".bak" not in d.name and any(is_state_dir(s) for s in d.iterdir()))


_STATE_CACHE = {}

# An independent copy of the documented, derivative-checked thermal contract.
# Do not import the production thermal reader or free-energy arithmetic here.
NATIVE_THERMAL_CONTRACT = dict(schema=1, model='solution geometry-level DFT Hessian', xc='r2scan',
    disp='d4', basis='def2-svp', anion_basis='def2-svpd', solvent='acetonitrile', solvent_model='SMD',
    density_fit=True, grid_level=3, cphf_grid_level=3, grid_response=True, conv_tol=1e-12,
    conv_tol_grad=1e-8, temperature_K=298.15, rotor_cutoff_cm=50,
    spin_partition='unsplit ground spin multiplet', gradient_max_acceptance_Ha_Bohr=.00045,
    min_frequency_acceptance_cm=-5, finite_difference_step_Bohr=.0005,
    finite_difference_relative_acceptance=.01, max_spin_contamination=.1, df_batch_cap_GiB=6,
    validated_pyscf_version='2.14.0', validated_gpu4pyscf_version='1.8.1')
NATIVE_CERTIFICATE = ROOT / 'config/native_thermal_validation.json'


def native_execution_budget_ok(record):
    """Independently verify an optional execution override, retaining legacy admission."""
    budget = record.get('execution_memory')
    if budget is None:
        return True
    try:
        if (budget.get('qualification_only') or budget.get('scientific_model_changed') is not False
                or budget.get('protocol_default_cap_GiB') != 6):
            return False
        if budget.get('cap_GiB') == 6:
            return budget.get('qualification_sha256') is None
        cap=budget.get('cap_GiB')
        if cap not in (16,24):
            return False
        path='config/native_hessian_memory_16_validation.json' if cap==16 else 'config/native_hessian_memory_validation.json'
        policy=('six-GiB integral batches; sixteen-GiB unrestricted response workspace' if cap==16
                else 'six-GiB integral batches; qualified larger unrestricted response workspace')
        certificate = json.loads((ROOT/path).read_text())
        seal = certificate.pop('content_sha256')
        if (budget.get('qualification_sha256') != seal
                or hashlib.sha256(json.dumps(certificate,sort_keys=True,separators=(',',':')).encode()).hexdigest() != seal
                or certificate.get('protocol') != NATIVE_THERMAL_CONTRACT
                or certificate.get('status') != 'execution_budget_qualified'
                or certificate.get('caps_GiB') != [cap]
                or certificate.get('scientific_model_changed') is not False
                or certificate.get('allocation_policy') != policy
                or budget.get('allocation_policy') != certificate.get('allocation_policy')
                or certificate.get('software') != record.get('software')):
            return False
        controls = certificate.get('controls', [])
        if len(controls) != 2 or {c.get('mult') for c in controls} != {1,2}:
            return False
        for control in controls:
            if not control.get('native_gates_passed'):
                return False
            for name in ('relative_matrix_error','maximum_matrix_difference','maximum_thermal_difference_eV'):
                value = fnum(control.get(name))
                if value is None or not 0 <= value <= 1e-6:
                    return False
            for name, digest in control['evidence_sha256'].items():
                if hashlib.sha256((ROOT/name).read_bytes()).hexdigest() != digest:
                    return False
        return True
    except (OSError,ValueError,KeyError,TypeError):
        return False



def native_execution_derivatives_ok(record,natm):
    """Independently reject unqualified numerical CDS changes or missing contact probes."""
    settings=record.get('execution_derivatives')
    if settings is None:return True
    try:
        step=settings.get('cds_step_Bohr')
        if step not in (1e-5,1e-6):return False
        filename='native_cds_step_validation.json' if step==1e-5 else 'native_cds_step_1e6_validation.json'
        certificate=json.loads((ROOT/'config'/filename).read_text())
        seal=certificate.pop('content_sha256')
        if (hashlib.sha256(json.dumps(certificate,sort_keys=True,separators=(',',':')).encode()).hexdigest()!=seal
                or settings.get('qualification_sha256')!=seal
                or certificate.get('status')!='execution_derivatives_qualified'
                or certificate.get('protocol')!=NATIVE_THERMAL_CONTRACT
                or certificate.get('cds_step_Bohr')!=step
                or certificate.get('library_default_cds_step_Bohr')!=1e-4
                or certificate.get('scientific_model_changed') is not False
                or certificate.get('software')!=record.get('software')
                or certificate.get('worst_contact_validation')!=dict(steps_Bohr=[.0005,.001],relative_acceptance=.01,symmetry_acceptance=1e-6)
                or settings.get('qualification_only') is not False
                or settings.get('scientific_model_changed') is not False
                or settings.get('source_geometry_sha256')!=record['identity']['geometry_sha256']):return False
        controls=certificate.get('controls',[])
        if len(controls)!=2 or {c.get('mult') for c in controls}!={1,2}:return False
        for control in controls:
            if control.get('native_gates_passed') is not True:return False
            for name in ('relative_matrix_error','maximum_matrix_difference','maximum_thermal_difference_eV'):
                value=fnum(control.get(name))
                if value is None or not 0<=value<=1e-6:return False
            evidence=control.get('evidence_sha256')
            if not isinstance(evidence,dict) or not evidence:return False
            for name,digest in evidence.items():
                if hashlib.sha256((ROOT/name).read_bytes()).hexdigest()!=digest:return False
        if step==1e-6:
            evidence=certificate.get('qualification_evidence_sha256')
            if certificate.get('target_native_gates_passed') is not True or not isinstance(evidence,dict) or not evidence:return False
            for name,digest in evidence.items():
                if hashlib.sha256((ROOT/name).read_bytes()).hexdigest()!=digest:return False
        indices=settings.get('worst_contact_indices')
        if (not isinstance(indices,list) or len(indices)!=2
                or any(type(i) is not int or not 0<=i<3*natm for i in indices)):return False
        checks=settings.get('worst_contact_checks',[])
        if not {(i,s) for i in indices for s in (.0005,.001)}<={(c.get('column'),c.get('step_Bohr')) for c in checks}:return False
        if any(fnum(c.get('relative_error')) is None or not 0<=c['relative_error']<=.01 for c in checks):return False
        evidence=settings.get('target_evidence_sha256')
        if not isinstance(evidence,dict) or not evidence:return False
        for name,digest in evidence.items():
            if hashlib.sha256((ROOT/name).read_bytes()).hexdigest()!=digest:return False
        return True
    except (OSError,ValueError,KeyError,TypeError,AttributeError):return False


def native_execution_pcm_ok(record):
    """Independently validate exact-tensor host memory execution provenance."""
    settings=record.get('execution_pcm_curvature')
    if settings is None:return True
    try:
        certificate=json.loads((ROOT/'config/native_pcm_curvature_memory_validation.json').read_text())
        seal=certificate.pop('content_sha256')
        if (hashlib.sha256(json.dumps(certificate,sort_keys=True,separators=(',',':')).encode()).hexdigest()!=seal
                or certificate.get('status')!='execution_pcm_curvature_qualified'
                or certificate.get('protocol')!=NATIVE_THERMAL_CONTRACT
                or certificate.get('software')!=record.get('software')
                or certificate.get('mode')!='host_analytic_tensors_atom_block_transfer'
                or certificate.get('scientific_model_changed') is not False or certificate.get('surface_grid_changed') is not False
                or settings.get('qualification_only') is not False or settings.get('production_admission') is not True
                or settings.get('scientific_model_changed') is not False or settings.get('surface_grid_changed') is not False
                or settings.get('mode')!=certificate.get('mode') or settings.get('qualification_sha256')!=seal
                or settings.get('source_geometry_sha256')!=record['identity']['geometry_sha256']):return False
        controls=certificate.get('controls',[])
        if len(controls)!=2 or {c.get('mult') for c in controls}!={1,2}:return False
        for c in controls:
            if c.get('native_gates_passed') is not True:return False
            for name in ('relative_matrix_error','maximum_matrix_difference','maximum_thermal_difference_eV'):
                value=fnum(c.get(name))
                if value is None or not 0<=value<=1e-6:return False
            evidence=c.get('evidence_sha256')
            if not isinstance(evidence,dict) or not evidence:return False
            for name,digest in evidence.items():
                if hashlib.sha256((ROOT/name).read_bytes()).hexdigest()!=digest:return False
        if hashlib.sha256((ROOT/certificate['isolated_proof_relative_path']).read_bytes()).hexdigest()!=certificate['isolated_proof_sha256']:return False
        for name in ('adapter_sha256','cpu_pcm_library_sha256','gpu_pcm_library_sha256','isolated_proof_sha256'):
            if settings.get(name)!=certificate.get(name):return False
        evidence=settings.get('target_evidence_sha256')
        if not isinstance(evidence,dict) or not evidence:return False
        for name,digest in evidence.items():
            if hashlib.sha256((ROOT/name).read_bytes()).hexdigest()!=digest:return False
        return True
    except (OSError,ValueError,KeyError,TypeError,AttributeError):return False

def _surface_thermal_inputs(sd, raw):
    """Independently admit the selected thermal record and recheck its arithmetic."""
    definition = 'legacy_gfn2_gas_rrho'
    policy = sd / 'thermal_policy.json'
    try:
        if policy.exists():
            selected = json.loads(policy.read_text())
            definition = selected['thermal_definition']
        if definition == 'legacy_gfn2_gas_rrho':
            return fnum(raw.get('g_thermal_eV')), raw.get('n_imag'), definition
        if definition != 'native_solution_dft_qrrho_spin':
            return None, None, definition
        if selected.get('native_protocol') != NATIVE_THERMAL_CONTRACT:
            return None, None, definition
        q, m = int(raw['charge']), int(raw['mult'])
        ident = dict(geometry_sha256=hashlib.sha256((sd/'opt.xyz').read_bytes()).hexdigest(),
                     charge=q, mult=m, protocol=NATIVE_THERMAL_CONTRACT)
        key = hashlib.sha256(json.dumps(ident, sort_keys=True).encode()).hexdigest()[:16]
        record = json.loads((sd/'native_thermal'/f'{key}.json').read_text())
        if record.get('identity') != ident or record.get('status') != 'derivative_checked_solution_thermal':
            return None, None, definition
        if not native_execution_budget_ok(record) or not native_execution_pcm_ok(record):
            return None, None, definition
        if any(record.get(k) is not True for k in ('converged','stationary_minimum','grid_response')):
            return None, None, definition
        lines = (sd/'opt.xyz').read_text().splitlines(); nat = int(lines[0])
        if not native_execution_derivatives_ok(record,nat):
            return None, None, definition
        symbols = [line.split()[0] for line in lines[2:2+nat]]
        coords = np.array([[float(x) for x in line.split()[1:4]] for line in lines[2:2+nat]])
        linear = np.linalg.matrix_rank(coords-coords.mean(axis=0), tol=1e-7) <= 1
        frequencies = record.get('frequencies_cm', [])
        if len(frequencies) != 3*nat-(5 if linear else 6) or any(fnum(v) is None for v in frequencies):
            return None, None, definition
        if min(frequencies) < -5:
            return None, None, definition
        for name, limit in (('gradient_max',.00045), ('maximum_asymmetry',1e-6)):
            value = fnum(record.get(name))
            if value is None or not 0 <= value <= limit:
                return None, None, definition
        checks = record.get('derivative_checks', [])
        oxygen = next((i for i,s in enumerate(symbols) if s=='O'),
                      next((i for i,s in enumerate(symbols) if s=='Fe'), nat-1))
        required = {(i,axis,step) for i,axis in ((0,1),(oxygen,2)) for step in (.0005,.001)}
        observed = {(c.get('atom_index'),c.get('axis'),c.get('step_Bohr')) for c in checks}
        if not required <= observed or any(fnum(c.get('relative_error')) is None or
                                           not 0 <= c['relative_error'] <= .01 for c in checks):
            return None, None, definition
        diagnostics = record.get('smd_spin_diagnostics', {})
        if any(fnum(diagnostics.get(k)) is None for k in ('spin_contam','homo_eV','s_squared')):
            return None, None, definition
        if abs(diagnostics['spin_contam']) > .1:
            return None, None, definition
        if record.get('software') != {'pyscf':'2.14.0','gpu4pyscf':'1.8.1'}:
            return None, None, definition
        proof = record.get('implementation_controls', [])
        expected = ['analytic_full_response_verified','uks_full_response_verified','fine_response_implementation_verified']
        if [p.get('status') for p in proof] != expected:
            return None, None, definition
        certificate = json.loads(NATIVE_CERTIFICATE.read_text())
        certificate_hash = certificate.pop('certificate_content_sha256')
        if (hashlib.sha256(json.dumps(certificate,sort_keys=True,separators=(',',':')).encode()).hexdigest()
                != certificate_hash or certificate.get('schema') != 1
                or certificate.get('native_protocol') != NATIVE_THERMAL_CONTRACT
                or certificate.get('software') != record['software']):
            return None, None, definition
        controls = certificate.get('controls', [])
        if [c.get('admission_status') for c in controls] != expected:
            return None, None, definition
        for evidence, control in zip(proof, controls):
            if evidence['sha256'] != control['source_sha256'] or control.get('software') != record['software']:
                return None, None, definition
            error = fnum(control.get('relative_full_matrix_error'))
            if error is None or not 0 <= error < .01:
                return None, None, definition
            if evidence['status'] == 'fine_response_implementation_verified':
                if (control.get('fine_response_implementation_verified') is not True
                        or evidence.get('relative_full_matrix_error') != error):
                    return None, None, definition
            path = Path(evidence['path'])
            if path.exists() and (hashlib.sha256(path.read_bytes()).hexdigest() != evidence['sha256']
                                 or json.loads(path.read_text()).get('status') != control['source_status']):
                return None, None, definition
        nuclear = fnum(record.get('g_nuclear_qrrho_eV'))
        spin = fnum(record.get('g_spin_partition_eV')); total = fnum(record.get('g_thermal_total_eV'))
        # k_B from the exact SI constants converted to eV; unsplit spin multiplicity.
        expected_spin = -(1.380649e-23/1.602176634e-19)*298.15*math.log(m)
        if None in (nuclear,spin,total) or abs(spin-expected_spin)>1e-8 or abs(total-nuclear-spin)>1e-8:
            return None, None, definition
        at50 = fnum((record.get('nuclear_by_cutoff_eV') or {}).get('50'))
        if at50 is None or abs(nuclear-at50)>1e-8:
            return None, None, definition
        return total, 0, definition
    except (OSError,ValueError,KeyError,TypeError):
        return None, None, definition


_PARTITION_AUDIT=None


def native_partition_inputs(sd,raw):
    """Use the independent internal-mode reader, never the producer loader."""
    global _PARTITION_AUDIT
    if _PARTITION_AUDIT is None:
        spec=importlib.util.spec_from_file_location('independent_native_partition',Path(__file__).with_name('native_partition.py'))
        _PARTITION_AUDIT=importlib.util.module_from_spec(spec);spec.loader.exec_module(_PARTITION_AUDIT)
    q,m=int(raw['charge']),int(raw['mult'])
    ident=dict(geometry_sha256=hashlib.sha256((sd/'opt.xyz').read_bytes()).hexdigest(),
               charge=q,mult=m,protocol=NATIVE_THERMAL_CONTRACT)
    address=hashlib.sha256(json.dumps(ident,sort_keys=True).encode()).hexdigest()[:16]
    return _PARTITION_AUDIT.read_partition(sd/'native_thermal'/(address+'.json'),sd/'opt.xyz')


def thermal_inputs(sd,raw):
    value,nimag,definition=_surface_thermal_inputs(sd,raw)
    if definition!='native_solution_dft_qrrho_spin' or value is None:
        return value,nimag,definition
    partition=native_partition_inputs(sd,raw)
    return (partition['g_thermal_total_eV'],0,definition) if partition is not None else (None,None,definition)


def spin_decision_ok(gid, st, raw):
    """Re-derive desired spin from the validated override or recorded UMA choice."""
    row = next((r for r in MANIFEST if r['id']==gid and r['state']==st), None)
    if row is None:
        return True
    try:
        override_path=ROOT/'config/spin_overrides.json'
        override=(json.loads(override_path.read_text()).get(gid,{}).get(st)
                  if override_path.exists() else None)
        if override is not None:
            if (override.get('protocol') != ACTIVE_SP or override.get('charge') != int(row['charge'])
                    or override.get('smiles_sha256') != hashlib.sha256(row['smiles'].encode()).hexdigest()):
                return False
            desired=int(override['mult'])
        else:
            uma=DFT.parent/'uma'/gid/st/'result.json'
            desired=(json.loads(uma.read_text()).get('chosen_mult') if uma.exists() else None)
            if desired is None:desired=int(row.get('mult_hint') or row.get('mult') or 1)
        return raw.get('mult') == desired
    except (OSError,ValueError,KeyError,TypeError):
        return False


def state(gid, st):
    """Raw state info: geometry/thermal from result.json + ACTIVE energy record."""
    k = (gid, st)
    if k in _STATE_CACHE:
        return _STATE_CACHE[k]
    sd = DFT / gid / st
    if not (sd / "result.json").exists():
        _STATE_CACHE[k] = None
        return None
    rj = json.loads((sd / "result.json").read_text())
    q, m = int(rj["charge"]), int(rj["mult"])
    rec = load_record(sd, q, m)                                  # converged only
    raw = load_record(sd, q, m, require_converged=False)          # for sanity
    g_th, n_imag, thermal_definition = thermal_inputs(sd, rj)
    d = dict(id=gid, state=st, charge=q, mult=m, dir=sd, rj=rj, rec=rec, raw_rec=raw,
             g_th=g_th, n_imag=n_imag, thermal_definition=thermal_definition,
             e_smd=(fnum(rec.get("e_smd_eV")) if rec else None),
             e_gas=(fnum(rec.get("e_gas_eV")) if rec else None),
             dG_solv=(fnum(rec.get("dG_solv_eV")) if rec else None))
    d["G"] = (d["e_smd"] + d["g_th"]) if (d["e_smd"] is not None and d["g_th"] is not None
                                                 and spin_decision_ok(gid,st,rj)) else None
    _STATE_CACHE[k] = d
    return d


def calc_states(gid):
    d = DFT / gid
    out = [state(gid, s.name) for s in d.iterdir() if is_state_dir(s)] if d.exists() else []
    return sorted([s for s in out if s], key=lambda s: -s["charge"])


def adjacent(states):
    return [(a, b) for a, b in zip(states, states[1:]) if a["charge"] - b["charge"] == 1]


def mtime_inputs(gid):
    """Newest mtime of result.json / opt.xyz / sp records / reorg caches of a molecule."""
    ts = []
    for p in (DFT / gid).rglob("*"):
        if p.is_file() and (p.suffix in (".json", ".xyz")):
            ts.append(p.stat().st_mtime)
    return max(ts) if ts else 0.0


ROWS = []


def add(table, gid, key, pub, rec, check_type="exact", digits=None, tol=TOL, note=""):
    """Record one comparison. Numbers are compared after rounding `rec` to `digits` (the
    published precision); strings/bools by equality."""
    def numlike(x):
        return (not isinstance(x, bool)) and (x is None or str(x).strip() == "" or fnum(x) is not None)
    if not (numlike(pub) and numlike(rec)):
        ps = "" if pub is None else str(pub)
        rs = "" if rec is None else str(rec)
        ok = ps.strip() == rs.strip()
        ROWS.append(dict(table=table, id=gid, key=key, published=ps, recomputed=rs,
                         abs_diff="", verdict="ok" if ok else "DISCREPANCY",
                         check_type=check_type, tol="equal", note=note))
        return ok
    p_num, r_num = fnum(pub), fnum(rec)
    if digits is not None and r_num is not None:
        r_cmp = round(r_num, digits)
    else:
        r_cmp = r_num
    if p_num is None and r_cmp is None:
        ok, diff = True, ""
    elif p_num is None or r_cmp is None:
        ok, diff = False, ""
    else:
        diff = abs(p_num - r_cmp)
        ok = diff <= tol
    ROWS.append(dict(table=table, id=gid, key=key,
                     published="" if p_num is None else repr(p_num),
                     recomputed="" if r_cmp is None else repr(r_cmp),
                     abs_diff="" if diff == "" else f"{diff:.3g}",
                     verdict="ok" if ok else "DISCREPANCY", check_type=check_type,
                     tol=f"{tol:g}", note=note))
    return ok


# ======================================================================== 1. E° (potentials)
def G_and_tq(s):
    """(G, missing_label) — G None if active energy or thermal missing (never 0)."""
    if s is None:
        return None, "no_energy"
    if s["e_smd"] is None:
        return None, "no_energy"
    if s["g_th"] is None:
        return None, "no_thermal"
    if s["G"] is None:
        return None, "wrong_spin"
    return s["G"], None


def fc_abs_live():
    o, n = state("ferrocene", "ox"), state("ferrocene", "neu")
    if (o is None or n is None or o["G"] is None or n["G"] is None
            or o['n_imag'] or n['n_imag'] or o['thermal_definition'] != n['thermal_definition']):
        return None
    return o["G"] - n["G"]


def manifest_groups():
    g = {}
    for r in MANIFEST:
        g.setdefault(r["id"], []).append(r)
    return g


MY_E = {}            # (id, event) -> dict(E_vs_Fc unrounded, status, ...)


def audit_potentials():
    T = "redox_potentials"
    fc = fc_abs_live()
    stored = ELE.FC_ABS_COMPUTED_V
    add(T, "ferrocene", "Fc_abs_live_vs_config_record", stored, fc, "cross_check", tol=0.005,
        note="config/project.json fc_abs_computed_V vs live G(Fc+)-G(Fc); pipeline warns at >5 mV")
    pub = {(r["id"], r["event"]): r for r in read_csv(RES / "redox_potentials.csv")}
    t_table = (RES / "redox_potentials.csv").stat().st_mtime
    seen = set()
    for gid, rows in manifest_groups().items():
        sts = sorted(((r["state"], int(r["charge"])) for r in rows), key=lambda x: -x[1])
        for (sO, qO), (sR, qR) in zip(sts, sts[1:]):
            if qO - qR != 1:
                continue
            ev = f"{sO}->{sR}"
            key = (gid, ev)
            seen.add(key)
            O, R = state(gid, sO), state(gid, sR)
            GO, mO = G_and_tq(O)
            GR, mR = G_and_tq(R)
            miss = [f"{s}({m})" for s, m in ((sO, mO), (sR, mR)) if m]
            status = "ok" if not miss else "INCOMPLETE:" + ",".join(miss)
            imag = any(x is not None and x["g_th"] is not None and (x["n_imag"] or 0) > 0
                       for x in (O, R))
            tqc = "imag_modes" if imag else ("ok" if not miss else "")
            if not miss and len({O['thermal_definition'], R['thermal_definition'],
                                  state('ferrocene','ox')['thermal_definition']}) != 1:
                tqc = 'thermal_definition_mismatch'
            unvalidated = not miss and tqc != 'ok'
            if unvalidated:status='THERMAL_UNVALIDATED'
            r = dict(status=status, thermal_qc=tqc)
            if not miss:
                r.update(dG=GR - GO, gO=O["g_th"], gR=R["g_th"], E_abs=GO - GR,
                         E=(GO - GR) - fc if fc is not None else None,
                         dEth=-(R["g_th"] - O["g_th"]))
                if unvalidated:
                    for quantity in ('dG','E_abs','E'):
                        r[quantity+'_estimate']=r[quantity];r[quantity]=None
            MY_E[key] = dict(r, q_ox=qO, q_red=qR, sO=sO, sR=sR)
            p = pub.get(key)
            stale = (" | molecule inputs newer than table" if mtime_inputs(gid) > t_table else "")
            if p is None:
                add(T, gid, f"{ev}:row_present", "", "present", "logic", note="couple missing from table")
                continue
            add(T, gid, f"{ev}:status", p.get("status"), status, "logic", note=stale.strip(" |"))
            add(T, gid, f"{ev}:thermal_qc", p.get("thermal_qc"), tqc, "logic")
            add(T, gid, f"{ev}:q_ox", p.get("q_ox"), qO, "logic", digits=0)
            add(T, gid, f"{ev}:dG_smd_eV", p.get("dG_smd_eV"), r.get("dG"), digits=4, note=stale.strip(" |"))
            add(T, gid, f"{ev}:g_thermal_ox_eV", p.get("g_thermal_ox_eV"), r.get("gO"), digits=4)
            add(T, gid, f"{ev}:g_thermal_red_eV", p.get("g_thermal_red_eV"), r.get("gR"), digits=4)
            add(T, gid, f"{ev}:E_abs_V", p.get("E_abs_V"), r.get("E_abs"), digits=3, note=stale.strip(" |"))
            add(T, gid, f"{ev}:E_vs_Fc_V", p.get("E_vs_Fc_V"), r.get("E"), digits=3, note=stale.strip(" |"))
            add(T, gid, f"{ev}:dE_thermal_V", p.get("dE_thermal_V"), r.get("dEth"), digits=4)
            for pubkey,quantity,digits in [('dG_smd_eV_estimate','dG_estimate',4),('E_abs_V_estimate','E_abs_estimate',3),
                                      ('E_vs_Fc_V_estimate','E_estimate',3)]:
                add(T,gid,f'{ev}:{pubkey}',p.get(pubkey),r.get(quantity),digits=digits)
            # Gas ML electronic differences are a separate diagnostic. They must use
            # the same model and a gas ML Fc reference, never the SMD/DFT free energy.
            ur=[]
            for ug,us in ((gid,sO),(gid,sR),('ferrocene','ox'),('ferrocene','neu')):
                up=DFT.parent/'uma'/ug/us/'result.json'
                ur.append(json.loads(up.read_text()) if up.exists() else None)
            valid=all(u and u.get('converged') is True and u.get('charge')==q
                      and fnum(u.get('energy_eV')) is not None for u,q in zip(ur,(qO,qR,1,0)))
            delta,value=None,None
            gas_status='INCOMPLETE_UMA_gas_diagnostic'
            if valid:
                models={u.get('model') for u in ur}
                if None in models or len(models)!=1:gas_status='model_mismatch'
                else:
                    delta=ur[1]['energy_eV']-ur[0]['energy_eV']
                    value=-delta-(ur[2]['energy_eV']-ur[3]['energy_eV'])
                    gas_status='UMA_gas_electronic_vs_UMA_Fc('+next(iter(models))+');not_MeCN_G'
            add(T,gid,f'{ev}:dE_gas_uma_eV',p.get('dE_gas_uma_eV'),delta,digits=4)
            add(T,gid,f'{ev}:E_vs_Fc_gas_V',p.get('E_vs_Fc_gas_V'),value,digits=3)
            add(T,gid,f'{ev}:gas_diagnostic_qc',p.get('gas_diagnostic_qc'),gas_status,'logic')
            if miss and fnum(p.get("E_vs_Fc_V")) is not None:
                add(T, gid, f"{ev}:incomplete_has_no_E", "E reported", "blank", "logic",
                    note="INCOMPLETE couple must not carry an E°")
            if (fnum(p.get("E_vs_Fc_V")) == 0.0 and gid != "ferrocene"
                    and (r.get("E") is None or abs(r["E"]) >= 5e-4)):   # 0 only if it rounds to 0
                add(T, gid, f"{ev}:E_is_exact_zero", "0", "nonzero", "logic", note="suspicious 0")
    for key in pub:
        if key not in seen:
            add(T, key[0], f"{key[1]}:row_present", "present", "", "logic",
                note="table row has no manifest couple")
    # couple sets: manifest (potentials) vs calcs/dft result.json charges (reorg/integrity/disp)
    for gid in sorted(manifest_groups()):
        computed = {x["state"] for x in calc_states(gid)}
        # a molecule still being computed publishes INCOMPLETE rows for its missing states;
        # compare only couples whose two states both exist
        man = {k[1] for k in MY_E if k[0] == gid
               and set(k[1].split("->")) <= computed}
        cal = {f"{a['state']}->{b['state']}" for a, b in adjacent(calc_states(gid))}
        if cal and man != cal:
            add(T, gid, "couple_set_manifest_vs_calcs", ";".join(sorted(man)), ";".join(sorted(cal)),
                "logic", note="potentials uses manifest states, integrity/reorg/disp use calcs dirs")
    # manifest charge/mult vs result.json
    for r in MANIFEST:
        s = state(r["id"], r["state"])
        if s is None:
            continue
        add("manifest", r["id"], f"{r['state']}:charge", r["charge"], s["charge"], "logic", digits=0)
        # the multiplicity is CHOSEN by the UMA spin scan (mult_hint is only a hint), so a
        # difference is not an arithmetic error — but UMA under-estimates spin gaps, so it is
        # a plausibility flag that calls for a DFT spin check
        if int(float(r["mult_hint"])) != int(s["mult"]):
            SAN.append(dict(id=r["id"], state=r["state"], check="mult_differs_from_hint",
                            value=f"{r['mult_hint']}->{s['mult']}", threshold="",
                            threshold_basis="UMA spin scan chose it; DFT-check (FINDINGS #22)",
                            note="multiplicity chosen by UMA differs from the manifest hint"))
    return fc


def audit_orop():
    """Independently reconstruct benchmark admission, including pending rows.

    Read the external experimental CSV and raw state/thermal evidence directly;
    do not import the production benchmark or its matching helper.
    """
    table = "orop_benchmark"
    dataset = ROOT / "data/raw/validation/OROP"
    with (dataset / "implicit_solvation_results.csv").open(encoding="utf-8-sig") as stream:
        source_rows = [{key.strip(): value for key, value in row.items()} for row in csv.DictReader(stream)]
    features = np.loadtxt(dataset / "features/feature_input_implicit_OROP_B3LYP-D3.csv",
                          skiprows=2, delimiter=",", ndmin=2)
    calc = ROOT / "calcs/orop"
    expected_systems = {int(path.name) for path in calc.iterdir() if path.is_dir() and path.name.isdigit()}
    published = {int(row["sys"]): row for row in read_csv(RES / "orop_benchmark.csv")}
    for system in sorted(expected_systems | set(published)):
        row = published.get(system)
        if system not in expected_systems or row is None:
            add(table, system, "row_present", "present" if row else "",
                "present" if system in expected_systems else "", "logic")
            continue
        charge = int(features[system - 1, 3])
        exp = float(source_rows[system - 1]["exp. redox pot. [V]"])
        original = float(source_rows[system - 1]["implicit solvation redox pot. [V]"])
        values, definitions, imaginary = [], [], []
        for name, wanted_charge in (("ox", charge), ("red", charge - 1)):
            directory = calc / str(system) / name
            energy = None; definition = None; nimag = None
            try:
                raw = json.loads((directory / "result.json").read_text())
                want = dict(opt_xc="r2scan", opt_basis="def2-svpd" if wanted_charge < 0 else "def2-svp",
                            opt_disp="d4", solvent="acetonitrile")
                geometry_ok = (raw.get("charge") == wanted_charge and raw.get("optimized") is True
                               and all(str(raw.get(key)).lower() == value for key, value in want.items()))
                record = load_record(directory, raw["charge"], raw["mult"]) if geometry_ok else None
                if record is not None:
                    correction, nimag, definition = thermal_inputs(directory, raw)
                    if correction is not None:
                        energy = record["e_smd_eV"] + correction
            except (OSError, ValueError, KeyError, TypeError):
                pass
            values.append(energy); definitions.append(definition); imaginary.append(nimag)
        fo, fn = state("ferrocene", "ox"), state("ferrocene", "neu")
        fc = fc_abs_live()
        reference_definition = fo['thermal_definition'] if fo else None
        all_definitions = definitions + [reference_definition, fn['thermal_definition'] if fn else None]
        complete = (all(value is not None for value in values) and fc is not None
                    and not any(imaginary) and fo is not None and fn is not None)
        status = "INCOMPLETE" if not complete else (
            "ok" if len(set(all_definitions)) == 1 else "thermal_definition_mismatch")
        potential = -(values[1] - values[0]) - fc if status == "ok" else None
        molecular_definition = definitions[0] if definitions[0] == definitions[1] else "mismatch"
        for key, expected, digits in (("exp", exp, None), ("charge_ox", charge, None),
                                      ("calc", potential, 3),
                                      ("err", potential - exp if potential is not None else None, 3),
                                      ("imp_dft", original, 3), ("err_impdft", original - exp, 3),
                                      ("reference_abs_V", fc, None)):
            add(table, system, key, row.get(key), expected, digits=digits)
        add(table, system, "status", row.get("status"), status, "logic")
        add(table, system, "thermal_definition", row.get("thermal_definition"), molecular_definition, "logic")
        add(table, system, "reference_thermal_definition", row.get("reference_thermal_definition"), reference_definition, "logic")
        add(table, system, "calibration_eligible", row.get("calibration_eligible"), str(status == "ok"), "logic")
        try:
            same_protocol = json.loads(row.get("energy_protocol", "")) == ACTIVE_SP
        except (ValueError, TypeError):
            same_protocol = False
        add(table, system, "energy_protocol", str(same_protocol), "True", "logic")


# =================================================================== 2. disproportionation
MY_DISP = {}   # (id, intermediate) -> dG eV


def audit_disp():
    T = "stability_disproportionation"
    pub = {(r["id"], r["intermediate"]): r for r in read_csv(RES / "stability_disproportionation.csv")}
    pot = {(r["id"], r["event"]): r for r in read_csv(RES / "redox_potentials.csv")}
    seen = set()
    for gid in calc_ids():
        sts = calc_states(gid)
        byq = {s["charge"]: s for s in sts}
        for s in sts:
            q = s["charge"]
            if (q + 1) in byq and (q - 1) in byq:
                hi, lo = byq[q + 1], byq[q - 1]
                key = (gid, s["state"])
                Gs = (hi["G"], s["G"], lo["G"])
                p = pub.get(key)
                if None in Gs:
                    if p is not None:
                        add(T, gid, f"{s['state']}:row_present", "present", "",
                            "logic", note="row published although an energy/thermal is missing")
                    continue
                seen.add(key)
                # E(hi/mid) - E(mid/lo), each E_abs = G(ox) - G(red)  (F*dE, 1 e-)
                E1 = hi["G"] - s["G"]
                E2 = s["G"] - lo["G"]
                dG = E1 - E2
                mismatch=len({x['thermal_definition'] for x in (hi,s,lo)})!=1
                valid=not any(x['n_imag'] for x in (hi,s,lo)) and not mismatch
                if valid:MY_DISP[key] = dG
                if p is None:
                    add(T, gid, f"{s['state']}:row_present", "", "present", "logic")
                    continue
                add(T,gid,f"{s['state']}:status",p.get('status'),'ok' if valid else 'THERMAL_UNVALIDATED','logic')
                add(T,gid,f"{s['state']}:thermal_qc",p.get('thermal_qc'),
                    'thermal_definition_mismatch' if mismatch else ('ok' if valid else 'imag_modes'),'logic')
                add(T, gid, f"{s['state']}:dG_disp_eV", p["dG_disp_eV"], dG if valid else None, digits=4)
                add(T, gid, f"{s['state']}:dG_disp_kJmol", p["dG_disp_kJmol"], dG * EV_KJMOL if valid else None, digits=1)
                add(T, gid, f"{s['state']}:logK_disp", p["logK_disp"],
                    -dG / (KT_EV_298 * math.log(10)) if valid else None, digits=2)
                add(T, gid, f"{s['state']}:stable_vs_disprop", p["stable_vs_disprop"], bool(dG > 0) if valid else '', "logic")
                add(T,gid,f"{s['state']}:dG_disp_eV_estimate",p.get('dG_disp_eV_estimate'),dG if not valid else None,digits=4)
                add(T,gid,f"{s['state']}:dG_disp_kJmol_estimate",p.get('dG_disp_kJmol_estimate'),dG*EV_KJMOL if not valid else None,digits=1)
                add(T,gid,f"{s['state']}:logK_disp_estimate",p.get('logK_disp_estimate'),
                    -dG/(KT_EV_298*math.log(10)) if not valid else None,digits=2)
                add(T,gid,f"{s['state']}:stable_vs_disprop_estimate",p.get('stable_vs_disprop_estimate'),
                    bool(dG>0) if not valid else '', 'logic')
                # cross-check vs the published E° table (3-dp rounding on each potential)
                e_hi = pot.get((gid, f"{hi['state']}->{s['state']}"))
                e_lo = pot.get((gid, f"{s['state']}->{lo['state']}"))
                if e_hi and e_lo and fnum(e_hi.get("E_abs_V")) is not None and fnum(e_lo.get("E_abs_V")) is not None:
                    add(T, gid, f"{s['state']}:dG_disp_vs_published_dE", p["dG_disp_eV"],
                        fnum(e_hi["E_abs_V"]) - fnum(e_lo["E_abs_V"]), "cross_check", tol=1.5e-3,
                        note="F*(E_hi - E_lo) from redox_potentials.csv (3-dp potentials)")
    for key in pub:
        if key not in seen and not any(r["key"] == f"{key[1]}:row_present" and r["id"] == key[0]
                                       for r in ROWS if r["table"] == T):
            add(T, key[0], f"{key[1]}:row_present", "present", "", "logic", note="unexpected row")
    # trace: stability_validation dG_calc vs our recompute
    for r in read_csv(RES / "stability_validation.csv"):
        cands = [v for (g, s), v in MY_DISP.items()
                 if g == r["id"] and s == r.get("intermediate", s)]
        if len(cands) == 1:
            add("stability_validation", r["id"], "dG_calc_kJmol", r["dG_calc_kJmol"],
                cands[0] * EV_KJMOL, "trace", digits=1)


# ============================================================================ 3. reorg (λ)
def kabsch(P, Q):
    P = P - P.mean(0); Q = Q - Q.mean(0)
    U, _, Vt = np.linalg.svd(P.T @ Q)
    d = np.sign(np.linalg.det(Vt.T @ U.T))
    R = Vt.T @ np.diag([1.0, 1.0, d]) @ U.T
    return float(np.sqrt(((P @ R.T - Q) ** 2).sum(1).mean()))


def read_xyz(p):
    lines = Path(p).read_text().splitlines()
    n = int(lines[0].split()[0])
    sym, xyz = [], []
    for ln in lines[2:2 + n]:
        t = ln.split()
        sym.append(t[0]); xyz.append([float(v) for v in t[1:4]])
    return sym, np.array(xyz)


def rmsd_naive(p1, p2):
    s1, x1 = read_xyz(p1); s2, x2 = read_xyz(p2)
    if s1 != s2:
        return None
    m = np.array(s1) != "H"
    return kabsch(x1[m], x2[m])


def rmsd_sym(p1, p2, charge):
    try:
        from rdkit import Chem
        from rdkit.Chem import rdDetermineBonds, rdMolAlign
        ms = []
        for p in (p1, p2):
            m = Chem.MolFromXYZFile(str(p))
            rdDetermineBonds.DetermineConnectivity(m, charge=int(charge))
            ms.append(Chem.RemoveHs(m, sanitize=False))
        return float(rdMolAlign.GetBestRMS(Chem.Mol(ms[1]), Chem.Mol(ms[0])))
    except Exception:
        return None


def cross_path(gid, species, at, q, m):
    g = geom_sha1(DFT / gid / at / "opt.xyz") or "-"
    return DFT / gid / "reorg" / f"{species}_at_{at}__{protocol_hash(g, q, m)}.json"


def cross_valid(d, q, m, ghash):
    """Validity of a cross-point cache under the active protocol (own re-statement)."""
    why = []
    if fnum(d.get("e_gas_eV")) is None: why.append("no finite e_gas")
    if d.get("protocol") != ACTIVE_SP: why.append("protocol provenance")
    if (d.get("sp_xc") or "").lower() != ACTIVE_SP["xc"]: why.append("xc")
    if (d.get("sp_basis") or "").lower() != ACTIVE_SP["basis"]: why.append("basis")
    if (d.get("sp_nlc") or None) != ACTIVE_SP["nlc"] or (d.get("sp_disp") or None) != ACTIVE_SP["disp"]:
        why.append("nlc/disp")
    if d.get("converged_gas") is not True: why.append("convergence unverified")
    if int(d.get("cache_schema", 1)) != 3: why.append("schema")
    if any(k not in d for k in ("gas_homo_eV", "anion_unbound")): why.append("diag")
    if int(d.get("charge", -999)) != q or int(d.get("mult", -999)) != m: why.append("charge/mult")
    if d.get("geom_sha1") != ghash: why.append("geom")
    return why


MY_LAM = {}   # (id, couple) -> dict


def audit_reorg():
    T = "reorganization"
    pub = {(r["id"], r["couple"]): r for r in read_csv(RES / "reorganization.csv")}
    t_table = (RES / "reorganization.csv").stat().st_mtime
    seen = set()
    for gid in calc_ids():
        for O, R in adjacent(calc_states(gid)):
            cp = f"{O['state']}->{R['state']}"
            key = (gid, cp)
            cO = cross_path(gid, O["state"], R["state"], O["charge"], O["mult"])  # E_O @ geom_R
            cR = cross_path(gid, R["state"], O["state"], R["charge"], R["mult"])  # E_R @ geom_O
            p = pub.get(key)
            problems = []
            if O["e_gas"] is None or R["e_gas"] is None:
                problems.append("own-geometry active record missing")
            dO = json.loads(cO.read_text()) if cO.exists() else None
            dR = json.loads(cR.read_text()) if cR.exists() else None
            if dO is None or dR is None:
                problems.append("cross-point cache missing")
            else:
                wO = cross_valid(dO, O["charge"], O["mult"], geom_sha1(DFT / gid / R["state"] / "opt.xyz"))
                wR = cross_valid(dR, R["charge"], R["mult"], geom_sha1(DFT / gid / O["state"] / "opt.xyz"))
                if wO or wR:
                    problems.append(f"cache invalid O:{wO} R:{wR}")
            if problems:
                MY_LAM[key] = None
                if p is not None:
                    add(T, gid, f"{cp}:row_present", "present", "", "logic", note="; ".join(problems))
                continue
            seen.add(key)
            unver = [c.name for c, d in ((cO, dO), (cR, dR)) if d.get("converged_gas") is None]
            if unver:
                add(T, gid, f"{cp}:crosspoint_convergence_recorded", "unverified(None)", "True", "logic",
                    note="cache accepted although converged_gas is absent: " + ",".join(unver)
                         + " (legacy-adopted; reorg.py _cache_ok rejects only False)")
            lam_O = dO["e_gas_eV"] - O["e_gas"]       # O surface: E_O(geom_R) - E_O(geom_O)
            lam_R = dR["e_gas_eV"] - R["e_gas"]       # R surface: E_R(geom_O) - E_R(geom_R)
            lam = lam_O + lam_R
            gO, gR = DFT / gid / O["state"] / "opt.xyz", DFT / gid / R["state"] / "opt.xyz"
            rn = rmsd_naive(gO, gR)
            rs = rmsd_sym(gO, gR, O["charge"])
            rm = rs if rs is not None else rn
            anion_records = [rec for q, rec in ((O["charge"], O["rec"]), (R["charge"], R["rec"]),
                                                (O["charge"], dO), (R["charge"], dR)) if q < 0]
            homos = [fnum(rec.get("gas_homo_eV")) for rec in anion_records]
            unb = any(rec.get("anion_unbound") is True for rec in anion_records) or any(
                h is not None and h > 0 for h in homos)
            why = []   # every failed check, '+'-joined (redox.properties.reorg contract)
            if lam < 0: why.append("negative_lambda")
            if min(lam_O, lam_R) < -0.02: why.append("negative_half")
            if rm is None: why.append("rmsd_unavailable")
            elif rm > 0.40: why.append("conformer_jump")
            if unb: why.append("anion_unbound")
            if any(h is None for h in homos): why.append("anion_diagnostics_missing")
            if max(lam_O, lam_R) > 1.0: why.append("large_half>1eV")
            flag = "+".join(why)
            MY_LAM[key] = dict(lam=lam, lam_O=lam_O, lam_R=lam_R, flag=flag, rmsd=rm)
            stale = "molecule inputs newer than table" if mtime_inputs(gid) > t_table else ""
            if p is None:
                add(T, gid, f"{cp}:row_present", "", "present", "logic", note=stale)
                continue
            add(T, gid, f"{cp}:lambda_i_eV", p["lambda_i_eV"], lam, digits=4, note=stale)
            add(T, gid, f"{cp}:lambda_i_meV", p["lambda_i_meV"], lam * 1000, digits=1)
            add(T, gid, f"{cp}:lambda_i_kJmol", p["lambda_i_kJmol"], lam * EV_KJMOL, digits=2)
            add(T, gid, f"{cp}:relax_ox_meV(lambda_O)", p["relax_ox_meV"], lam_O * 1000, digits=3)
            add(T, gid, f"{cp}:relax_red_meV(lambda_R)", p["relax_red_meV"], lam_R * 1000, digits=3)
            add(T, gid, f"{cp}:rmsd_naive_A", p["rmsd_naive_A"], rn, digits=3)
            add(T, gid, f"{cp}:rmsd_A", p["rmsd_A"], rm, digits=3)
            add(T, gid, f"{cp}:anion_homo_eV", p["anion_homo_eV"],
                max((h for h in homos if h is not None), default=None), digits=3)
            add(T, gid, f"{cp}:flag", p["flag"], flag, "logic", note=stale)
            # QC gap check: the own-geometry ANION record can itself be gas-unbound (HOMO>0)
            own_homo = fnum(R["rec"].get("gas_homo_eV"))
            if R["charge"] < 0 and own_homo is not None and own_homo > 0 and flag == "":
                add(T, gid, f"{cp}:own_anion_unbound_not_flagged", "", f"gas HOMO(R@R)={own_homo:.3f} eV",
                    "logic", note="flag only inspects the cross point E_R@geom_O, not E_R@geom_R")
    for key in pub:
        if key not in seen and not any(r["id"] == key[0] and r["key"] == f"{key[1]}:row_present"
                                       for r in ROWS if r["table"] == T):
            add(T, key[0], f"{key[1]}:row_present", "present", "", "logic", note="unexpected row")


# ========================================================================= 6. integrity
MY_INTEG = {}


def audit_integrity():
    T = "state_integrity"
    pub = {(r["id"], r["couple"]): r for r in read_csv(RES / "state_integrity.csv")}
    seen = set()
    for gid in calc_ids():
        for O, R in adjacent(calc_states(gid)):
            cp = f"{O['state']}->{R['state']}"
            key = (gid, cp)
            seen.add(key)
            p = pub.get(key)
            if O["e_smd"] is None or R["e_smd"] is None:
                ea, bound = None, None
            else:
                ea = O["e_smd"] - R["e_smd"]
                bound = not (R["charge"] < 0 and ea <= 0.0)
            db = fnum(p.get("d_bonds")) if p else None      # bond graph: taken from table
            if bound is None:
                verdict = "incomplete"
            elif not bound:
                verdict = "unbound"
            elif db is None or db > 0:
                verdict = "bond_change"
            else:
                verdict = "intact_bound"
            MY_INTEG[key] = verdict
            if p is None:
                add(T, gid, f"{cp}:row_present", "", "present", "logic")
                continue
            add(T, gid, f"{cp}:EA_solv_eV", p["EA_solv_eV"], ea, digits=3)
            add(T, gid, f"{cp}:verdict", p["verdict"], verdict, "logic",
                note="bound part recomputed; d_bonds taken from the table")
    for key in pub:
        if key not in seen:
            add(T, key[0], f"{key[1]}:row_present", "present", "", "logic", note="unexpected row")


# ============================================================================ 4. capacity
TOLYL = "[CH3]c1ccc([CH2])cc1"


def mw_from_smiles(smi):
    from rdkit import Chem
    from rdkit.Chem import Descriptors
    m = Chem.MolFromSmiles(smi)
    return (Descriptors.MolWt(m), len(m.GetSubstructMatches(Chem.MolFromSmarts(TOLYL)))) if m else (None, 0)


def mw_from_xyz(p):
    from rdkit import Chem
    pt = Chem.GetPeriodicTable()
    sym, _ = read_xyz(p)
    return sum(pt.GetAtomicWeight(s) for s in sym)


def scorecard_masses(smi):
    """Reconstruct the documented four-decimal mass inputs to the scorecard.

    The capacity table rounds masses to four decimals before serialization.
    Applying a second display rounding directly to the raw floating sum can
    cross a decimal tie (402.4500000000002 versus serialized 402.45).
    Independently reproduce that interface without reading its mass values.
    """
    mass, sites = mw_from_smiles(smi)
    if mass is None:
        return None, None
    return round(mass, 4), round(mass + sites * C_MASS, 4) if sites else None


MY_MW = {}


def audit_capacity():
    T = "capacity_and_proxies"
    smi = {}
    for r in MANIFEST:
        if r.get("smiles"):
            smi.setdefault(r["id"], r["smiles"])
    pub = {r["id"]: r for r in read_csv(RES / "capacity_and_proxies.csv")}
    rest = {r["id"]: r["state"] for r in MANIFEST if str(r.get("n_e")) == "0"}
    for gid, p in pub.items():
        s = smi.get(gid)
        mw, ng = mw_from_smiles(s) if s else (None, 0)
        MY_MW[gid] = (mw, ng)
        sts = calc_states(gid)
        qs = sorted({int(row["charge"]) for row in MANIFEST if row["id"] == gid})
        n_all = sum(1 for a, b in zip(qs, qs[1:]) if b - a == 1)
        add(T, gid, "n_electrons(all adjacent couples)", p["n_electrons"], n_all, digits=0)
        add(T, gid, "MW", p["MW"], mw, digits=4)
        add(T, gid, "n_graft_sites", p["n_graft_sites"], ng, digits=0)
        add(T, gid, "MW_repeat_unit", p["MW_repeat_unit"], (mw + ng * C_MASS) if ng else None, digits=4)
        add(T, gid, "specific_capacity_mAh_g", p["specific_capacity_mAh_g"],
            n_all * F_TABLE_CAPACITY / (mw * C_PER_MAH) if (mw and n_all) else None, digits=1)
        neu = [x for x in sts if x["charge"] == 0]
        dgs = next((x["dG_solv"] for x in neu if x["dG_solv"] is not None), None)
        add(T, gid, "dGsolv_neutral_eV", p["dGsolv_neutral_eV"], dgs, digits=3)
        # SMILES formula vs the stored resting-state geometry (H count / atoms)
        rs = rest.get(gid)
        if rs and (DFT / gid / rs / "opt.xyz").exists() and mw:
            add(T, gid, "MW_smiles_vs_geometry", mw, mw_from_xyz(DFT / gid / rs / "opt.xyz"),
                "cross_check", tol=0.01, note=f"RDKit MolWt(SMILES) vs atoms in {rs}/opt.xyz")


# =========================================================== 5a. scorecard (independent)
def candidate_meta():
    batches, flags, unrank = {}, {}, set()
    for b, mod in (("starting", "starting_candidates"), ("merrifield_multi", "merrifield_multielectron"),
                   ("discovered", "discovered_candidates"), ("custom", "custom_candidates")):
        for g in cfg(mod).GROUPS:
            batches[g["id"]] = b
            if g.get("rankable") is False:
                unrank.add(g["id"])
    batches["viologen"] = "starting"
    return batches, unrank


def ion_mass(smiles):
    from rdkit import Chem
    from rdkit.Chem import Descriptors
    return Descriptors.MolWt(Chem.MolFromSmiles(smiles))


def rmse(xs):
    xs = [x for x in xs if x is not None]
    return math.sqrt(sum(x * x for x in xs) / len(xs)) if xs else None


def robust_sigma(xs):
    xs = sorted(xs)
    if not xs:
        return None
    med = statistics.median(xs)
    return 1.4826 * statistics.median([abs(x - med) for x in xs])


def my_sigmas(fc):
    # Grounded experimental points, one per (molecule, couple): tier-A E events of
    # config/benchmark_mecn.py (several sources averaged), then config/validation.py events
    # with grounded=True that the benchmark does not cover; the Fc reference never counts.
    famap = (("quinone", "quinone (n-type)"), ("aromatic imide", "imide (n-type)"),
             ("viologen", "pyridine-multi-e"), ("phenothiazine", "amine (p-type)"),
             ("nitroxide", "nitroxide"))
    pts = {}
    for b in cfg("benchmark_mecn").BENCHMARK:
        if b.get("excluded"):
            continue
        f = next((v for k, v in famap if b["bench_family"].startswith(k)), None)
        for ev in b["events"]:
            if ev["kind"] == "E" and ev["tier"] == "A":
                pts.setdefault((b["id"], ev["event"]), [f, []])[1].append(ev["exp_V_vs_Fc"])
    for e in cfg("validation").VALIDATION:
        if e["id"] == "ferrocene":
            continue
        for ev in e.get("events", []):
            k = (e["id"], ev["event"])
            if ev.get("grounded") and k not in pts:
                pts[k] = [e.get("family"), [ev["exp_V_vs_Fc"]]]
    fam, allres = {}, []
    for k, (f, vals) in pts.items():
        r = MY_E.get(k)
        if not r or r.get("E") is None:
            continue
        res = round(r["E"], 3) - sum(vals) / len(vals)
        allres.append(res)
        if f:
            fam.setdefault(f, []).append(res)
    minimum = SCC.MIN_REDOX_CALIBRATION_POINTS
    fam_s = {f: round(rmse(v), 3) for f, v in fam.items() if len(v) >= minimum}
    pooled = round(rmse(allres), 3) if len(allres) >= minimum else None
    lam = []
    for r in read_csv(RES / "d3tales_reorg_validation_prod" / "comparison.csv"):
        a, b = fnum(r.get("our_smd_eV")), fnum(r.get("d3tales_eV"))
        if a is not None and b is not None and 0 <= a <= 2 and 0 <= b <= 2:
            lam.append(a - b)
    s_lam = round(robust_sigma(lam), 3) if lam else SCC.SIGMA_LAMBDA_EV
    dres = [fnum(r.get("err_kJmol")) for r in read_csv(RES / "stability_validation.csv")
            if (r.get("tier") or "A") == "A"]
    dres = [x / EV_KJMOL for x in dres if x is not None]
    s_disp = round(rmse(dres), 3) if dres else SCC.SIGMA_DISP_EV
    return fam_s, pooled, s_lam, s_disp


def my_spin_evidence(gid):
    """Rebuild source identities and gaps independently; pending helpers never rank."""
    import fcntl
    from redox.core.protocol import record_path
    lock=ROOT/'logs/pipeline_locks'/f'{gid}_spincheck.lock'
    try:
        with lock.open('r') as stream:
            try:fcntl.flock(stream,fcntl.LOCK_SH|fcntl.LOCK_NB)
            except BlockingIOError:return None,'pending:active_spin_comparison','DFT_check_required'
    except FileNotFoundError:pass
    gaps, sources, pending = {}, {}, {}
    digest = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
    def native_path(directory, charge, mult):
        identity = dict(geometry_sha256=digest(directory/'opt.xyz'), charge=charge,
                        mult=mult, protocol=NATIVE_THERMAL_CONTRACT)
        key = hashlib.sha256(json.dumps(identity, sort_keys=True).encode()).hexdigest()[:16]
        return directory/'native_thermal'/f'{key}.json'
    for path in (DFT.parent/'uma'/gid).glob('*/result.json'):
        raw=json.loads(path.read_text());gap=fnum(raw.get('spin_gap_eV'))
        if gap is not None:gaps[path.parent.name]=abs(gap);sources[path.parent.name]='uma'
    for path in (DFT.parent/'spincheck'/gid).glob('*_m*/result.json'):
        st=path.parent.name.rsplit('_m',1)[0];alt=json.loads(path.read_text());chosen=state(gid,st)
        if chosen is None or alt.get('charge')!=chosen['charge'] or alt.get('mult')==chosen['mult']:continue
        sd=chosen['dir'];row=next((r for r in MANIFEST if r['id']==gid and r['state']==st),None)
        if row is None or chosen['rec'] is None:
            pending[st]='source_addressed_spin_report_pending';continue
        identity=dict(id=gid,state=st,charge=chosen['charge'],mult=chosen['mult'],
            geometry_sha256=digest(sd/'opt.xyz'),smiles_sha256=hashlib.sha256(row['smiles'].encode()).hexdigest(),
            sp_record_sha256=digest(record_path(sd,chosen['charge'],chosen['mult'])),protocol=ACTIVE_SP)
        address=hashlib.sha256(json.dumps(identity,sort_keys=True).encode()).hexdigest()[:16]
        report_path=sd/'spin_validation'/f'{address}.json'
        report=json.loads(report_path.read_text()) if report_path.exists() else None
        if not report or report.get('chosen_identity')!=identity or report.get('status') not in ('spin_gap_verified','lower_spin_promoted'):
            pending[st]='source_addressed_spin_report_pending';continue
        if alt.get('spin_validation_status') not in (None,'validated'):
            pending[st]='alternate_spin_validation_pending';continue
        if report.get('ensemble_primary_integration_pending'):
            pending[st]='spin_ensemble_integration_pending';continue
        expected=dict(opt_xc='r2scan',opt_basis='def2-svpd' if alt['charge']<0 else 'def2-svp',
                      opt_disp='d4',solvent=ELE.SOLVENT['name'])
        if (alt['mult']!=report.get('alternative_mult') or not alt.get('optimized') or
                any(str(alt.get(k)).lower()!=str(v).lower() for k,v in expected.items())):
            pending[st]='alternate_protocol_or_multiplicity_mismatch';continue
        record=load_record(path.parent,int(alt['charge']),int(alt['mult']))
        if record is None or chosen['e_smd'] is None:
            pending[st]='alternate_active_SP_pending';continue
        gap=record['e_smd_eV']-chosen['e_smd'];reported=fnum(report.get('electronic_gap_eV'))
        contamination=fnum(record.get('smd_spin_contam'))
        if reported is None or abs(gap-reported)>1e-8 or (contamination is not None and abs(contamination)>.1):
            pending[st]='alternate_energy_or_spin_report_stale';continue
        gt, ni, definition=thermal_inputs(path.parent,alt)
        matching=(definition==chosen['thermal_definition']=='native_solution_dft_qrrho_spin')
        if report.get('thermochemical_gap_validated'):
            selected_partition=native_partition_inputs(sd,chosen['rj'])
            alternate_partition=native_partition_inputs(path.parent,alt)
            partition_dependencies=report.get('native_partition_dependencies') or {}
            if (selected_partition is None or alternate_partition is None
                    or report.get('native_partition_protocol')!=selected_partition['protocol']
                    or alternate_partition['protocol']!=selected_partition['protocol']
                    or partition_dependencies.get('chosen')!=selected_partition['sha256']
                    or partition_dependencies.get('alternative')!=alternate_partition['sha256']):
                pending[st]='native_spin_partition_dependencies_stale';continue
            cp=native_path(sd,chosen['charge'],chosen['mult']);ap=native_path(path.parent,alt['charge'],alt['mult'])
            deps=report.get('native_dependencies') or {}
            if not cp.exists() or not ap.exists() or deps.get('chosen')!=digest(cp) or deps.get('alternative')!=digest(ap):
                pending[st]='native_spin_dependencies_stale';continue
        native=(matching and gt is not None and chosen['G'] is not None and ni==0)
        if native:
            gap+=gt-chosen['g_th'];reported=fnum(report.get('thermochemical_gap_eV'))
            if not report.get('thermochemical_gap_validated') or reported is None or abs(gap-reported)>1e-8:
                pending[st]='native_spin_report_pending_or_stale';continue
            source='dft_free_energy'
        else:source='dft_electronic(thermal_unvalidated)'
        if native and abs(gap)<=.1:pending[st]='spin_ensemble_required'
        elif not native and abs(gap)<=.2:pending[st]='matching_native_spin_thermals_required'
        gaps[st]=gap;sources[st]=source
    if pending:
        st=min(pending);return None,f'pending:{pending[st]}:{st}','DFT_check_required'
    if not gaps:return None,'','ok'
    st=min(gaps,key=gaps.get);value=round(gaps[st],4);source=sources[st]+':'+st
    confidence=('wrong_ground_state' if source.startswith('dft_') and value<-1e-4 else
                ('DFT_check_required' if source.startswith('uma:') and value<.5 else
                 ('low' if value<SCC.SPIN_GAP_LOW_EV else 'ok')))
    return value,source,confidence


def scorecard_inner_lambdas(raw):
    """Reconstruct serialized per-couple inputs, then the scorecard's mean."""
    if not raw:
        return None, None, None
    return tuple(round(statistics.mean(values), 4) for values in (
        [round(r["lam"], 4) for r in raw],
        [round(r["lam_O"] * 1000, 3) / 1000 for r in raw],
        [round(r["lam_R"] * 1000, 3) / 1000 for r in raw],
    ))


def my_scorecard(fc):
    batches, unrank = candidate_meta()
    wlo, whi = ELE.WINDOW_V_VS_FC
    div = ELE.ANOLYTE_CATHOLYTE_DIVIDER_V
    fam_s, pooled, s_lam, s_disp = my_sigmas(fc)
    # Independently reconstruct the current charge-class fallback rather than
    # borrowing the production calibration admission or sigma selector.
    errors = {"cation": [], "anion": [], "multi": []}
    reference_definition = state("ferrocene", "ox")["thermal_definition"]
    for row in read_csv(RES / "orop_benchmark.csv"):
        try:
            eligible = (row.get("status") == "ok" and row.get("calibration_eligible") == "True"
                        and json.loads(row["energy_protocol"]) == ACTIVE_SP
                        and row.get("thermal_definition") == row.get("reference_thermal_definition") == reference_definition
                        and fc is not None and abs(float(row["reference_abs_V"]) - fc) < 1e-8)
            error, charge = fnum(row.get("err")), fnum(row.get("charge_ox"))
        except (ValueError, TypeError, KeyError):
            eligible = False; error = charge = None
        if eligible and error is not None and charge is not None and fnum(row.get("calc")) is not None:
            kind = "multi" if charge >= 2 or charge < 0 else ("cation" if charge >= 1 else "anion")
            errors[kind].append(error)
    charge_prior = {kind: round(rmse(values), 3) if len(values) >= SCC.MIN_REDOX_CALIBRATION_POINTS
                    else SCC.SIGMA_REDOX_V[kind] for kind, values in errors.items()}
    rest = {r["id"]: (r["state"], int(r["charge"])) for r in MANIFEST if str(r.get("n_e")) == "0"}
    lam_o = {}
    for r in read_csv(RES / "lambda_outer.csv"):
        st = frozenset(x.strip() for x in r["couple"].replace("->", "/").split("/"))
        lam_o[(r["id"], st)] = (fnum(r["lambda_o_pcm_eV"]), r.get("method", ""))
    cap_tab = {r["id"]: r for r in read_csv(RES / "capacity_and_proxies.csv")}
    m_an = ion_mass(SCC.COUNTERION_ANION_SMILES)
    m_cat = {k: ion_mass(s) for k, s in SCC.SUPPORTING_CATIONS.items()}
    pekar = 1 / ELE.SOLVENT["eps_optical"] - 1 / ELE.SOLVENT["eps_r"]
    fam_of = {r["id"]: r["family"] for r in MANIFEST}
    out = {}
    for gid in sorted(batches):
        if gid == "ferrocene":
            continue
        couples = {}
        for (g, ev), r in MY_E.items():
            if g != gid:
                continue
            couples[r["q_ox"]] = dict(r, couple=ev, integ=MY_INTEG.get((gid, ev)),
                                      lam=MY_LAM.get((gid, ev)),
                                      lo=lam_o.get((gid, frozenset((r["sO"], r["sR"])))))
        if not couples:
            continue
        if gid not in rest:
            out[(gid, "")] = dict(status="INCOMPLETE", n_accessible=0)
            continue
        q0 = rest[gid][1]
        paths, stop = {}, {}
        for pool, step in (("anolyte", -1), ("catholyte", 1)):
            acc, q = [], q0
            while True:
                c = couples.get(q if step < 0 else q + 1)
                if c is None:
                    stop[pool] = "end of computed ladder"; break
                Er = round(c["E"], 3) if c.get("E") is not None else None
                if c["status"] != "ok" or Er is None:
                    stop[pool] = f"INCOMPLETE {c['couple']}: {c['status']}"; break
                if c["integ"] in (None, "incomplete"):
                    stop[pool] = f"INCOMPLETE {c['couple']}: integrity not assessed"; break
                inwin = wlo <= Er <= whi
                side = (Er < div) if pool == "anolyte" else (Er >= div)
                if not (c["integ"] == "intact_bound" and inwin and side):
                    why = ("integrity=" + c["integ"] if c["integ"] != "intact_bound"
                           else ("outside window" if not inwin else "other side of divider"))
                    stop[pool] = f"{c['couple']} inaccessible ({why})"; break
                acc.append(dict(c, Er=Er)); q += step
            paths[pool] = acc
        if not paths["anolyte"] and not paths["catholyte"]:
            out[(gid, "")] = dict(status=("INCOMPLETE" if any(v.startswith("INCOMPLETE") for v in stop.values())
                                          else "REJECTED"), n_accessible=0)
            continue
        mw, mw_rep = scorecard_masses(next(r["smiles"] for r in MANIFEST if r["id"] == gid))
        fam = fam_of.get(gid)
        sigE = fam_s.get(fam, pooled)
        spin_gap,spin_source,spin_confidence=my_spin_evidence(gid)
        for pool, acc in paths.items():
            if not acc:
                continue
            pool_sigma = sigE
            if pool_sigma is None:
                pool_sigma = max(charge_prior["multi" if max(abs(c["q_ox"]), abs(c["q_red"])) >= 2
                                               else ("anion" if min(c["q_ox"], c["q_red"]) < 0 else "cation")]
                                 for c in acc)
            n = len(acc)
            capf = lambda mass: n * FARADAY / (mass * C_PER_MAH) if mass else None  # noqa: E731
            qs = {q for c in acc for q in (c["q_ox"], c["q_red"])}
            def maxload(mc):
                return max([q * m_an if q > 0 else (-q * mc if q < 0 else 0.0) for q in qs])
            clean = [c for c in acc if c["lam"] and not c["lam"]["flag"]]
            flagged = [c for c in acc if c["lam"] and c["lam"]["flag"]]
            missing = [c for c in acc if not c["lam"]]
            mean = lambda xs: statistics.mean(xs) if xs else None  # noqa: E731
            # The scorecard consumes the independently audited reorganization
            # CSV: totals are serialized to 4 decimals in eV, halves to 3 in meV.
            # Reconstruct those published inputs from raw energies before its
            # mean, rather than compare against a differently rounded raw mean.
            lam4, lO, lR = scorecard_inner_lambdas([c["lam"] for c in clean])
            los = [c["lo"][0] for c in clean if c["lo"] and c["lo"][0]]
            lo1 = round(mean(los), 4) if los else None
            lam_het = lam4 / 2 + lo1 if (lam4 is not None and lo1) else None
            lam_se = lam4 + lo1 if (lam4 is not None and lo1) else None  # 2*lo1*(1-a/2a)
            if pool == "anolyte":
                inter = [c["sR"] for c in acc[:-1]]
            else:
                inter = [c["sO"] for c in acc[:-1]]
            dv = [MY_DISP[(gid, s)] * EV_KJMOL for s in inter if (gid, s) in MY_DISP]
            ct = cap_tab.get(gid, {})
            out[(gid, pool)] = dict(
                # missing-data truncation -> INCOMPLETE (only a lower bound), else candidate
                status=("INCOMPLETE" if str(stop[pool]).startswith("INCOMPLETE") else "candidate"),
                min_spin_gap_eV=spin_gap,spin_gap_source=spin_source,spin_confidence=spin_confidence,
                rankable=(gid not in unrank), resting_state=rest[gid][0],
                path=" ; ".join(c["couple"] for c in acc), path_stop=stop[pool], n_accessible=n,
                E_V=statistics.mean(c["Er"] for c in acc),
                sigma_E_V=pool_sigma,
                specific_capacity_mAh_g=capf(mw), MW=mw, MW_repeat_unit=mw_rep,
                capacity_repeat_mAh_g=capf(mw_rep),
                capacity_maxload_Li_mAh_g=capf(mw_rep + maxload(m_cat["Li"])) if mw_rep else None,
                capacity_maxload_TBA_mAh_g=capf(mw_rep + maxload(m_cat["TBA"])) if mw_rep else None,
                lambda_i_eV=lam4, lambda_i_ox_eV=lO, lambda_i_red_eV=lR,
                sigma_lambda_eV=(s_lam if lam4 is not None else None),
                lambda_qc=("ok" if len(clean) == n else ("all_flagged" if not clean else "partial")),
                n_lambda_flagged=len(flagged), n_lambda_missing=len(missing),
                lambda_flags=";".join(sorted({c["lam"]["flag"] for c in flagged})),
                lambda_o_pcm_eV=lo1, lambda_het_eV=lam_het, lambda_se_contact_eV=lam_se,
                sigma_lambda_het_eV=(math.sqrt((s_lam / 2) ** 2 + SCC.SIGMA_LAMBDA_O_EV ** 2)
                                     if lam_het is not None else None),
                dG_disp_kJmol=(min(dv) if dv else None), disp_applicable=(n >= 2),
                sigma_disp_eV=(s_disp if dv else None),
                thermal_qc=("imag_modes" if any(c["thermal_qc"] == "imag_modes" for c in acc) else "ok"),
                SA_score=fnum(ct.get("SA_score")),
                dGsolv_proxy_eV=next((x["dG_solv"] for x in calc_states(gid) if x["charge"] == 0), None),
            )
    return out


SC_FIELDS = [  # (column, digits, check_type)
    ('min_spin_gap_eV',4,'exact'),('spin_gap_source',None,'logic'),('spin_confidence',None,'logic'),
    ("status", None, "logic"), ("resting_state", None, "logic"), ("path", None, "logic"),
    ("path_stop", None, "logic"), ("n_accessible", 0, "logic"), ("E_V", 3, "exact"),
    ("sigma_E_V", 3, "exact"), ("specific_capacity_mAh_g", 1, "exact"), ("MW", 1, "exact"),
    ("MW_repeat_unit", 1, "exact"), ("capacity_repeat_mAh_g", 1, "exact"),
    ("capacity_maxload_Li_mAh_g", 1, "exact"), ("capacity_maxload_TBA_mAh_g", 1, "exact"),
    ("lambda_i_eV", 4, "exact"), ("lambda_i_ox_eV", 4, "exact"), ("lambda_i_red_eV", 4, "exact"),
    ("sigma_lambda_eV", 3, "exact"), ("lambda_qc", None, "logic"), ("n_lambda_flagged", 0, "logic"),
    ("n_lambda_missing", 0, "logic"), ("lambda_flags", None, "logic"),
    ("lambda_o_pcm_eV", 4, "trace"), ("lambda_het_eV", 4, "exact"),
    ("sigma_lambda_het_eV", 3, "exact"), ("lambda_se_contact_eV", 4, "exact"),
    ("dG_disp_kJmol", 1, "exact"), ("disp_applicable", None, "logic"),
    ("sigma_disp_eV", 3, "exact"), ("thermal_qc", None, "logic"), ("SA_score", 2, "trace"),
    ("dGsolv_proxy_eV", 3, "exact"),
]


def audit_scorecard(fc):
    T = "scorecard"
    mine = my_scorecard(fc)
    pub = {}
    for r in read_csv(RES / "scorecard.csv"):
        pub[(r["id"], r.get("pool") or "")] = r
    for key in sorted(set(mine) | set(pub)):
        gid, pool = key
        m, p = mine.get(key), pub.get(key)
        tag = f"{pool or '-'}"
        if m is None or p is None:
            add(T, gid, f"{tag}:row_present", "present" if p else "", "present" if m else "", "logic")
            continue
        for col, dg, ct in SC_FIELDS:
            mv = m.get(col)
            if isinstance(mv, bool):
                add(T, gid, f"{tag}:{col}", p.get(col), str(mv), ct)
            elif dg is None or isinstance(mv, str):
                add(T, gid, f"{tag}:{col}", p.get(col), "" if mv is None else mv, ct)
            else:
                add(T, gid, f"{tag}:{col}", p.get(col), mv, ct, digits=dg)
        if p.get("status") == "candidate" and str(p.get("path_stop", "")).startswith("INCOMPLETE"):
            add(T, gid, f"{tag}:candidate_with_incomplete_path", "candidate", "INCOMPLETE?", "logic",
                note="path truncated by missing data, yet row is a candidate (capacity is a lower bound)")
    return mine


# ============================================================== 5b. Pareto (independent)
def objectives(row, pool, cap_key):
    o, miss = {}, []
    if row.get('spin_confidence') in ('DFT_check_required','wrong_ground_state'):miss.append('spin_state')
    E = fnum(row.get("E_V"))
    thermal_valid = row.get("thermal_qc", "ok") == "ok"
    if E is None or not thermal_valid: miss.append("voltage")
    else: o["voltage"] = (-E if pool == "anolyte" else E, fnum(row.get("sigma_E_V")) or 0.0)
    cap = fnum(row.get(cap_key))
    if cap is None: miss.append("capacity")
    else: o["capacity"] = (cap, 0.0)
    dg = fnum(row.get("dG_disp_kJmol"))
    if dg is not None and thermal_valid:
        o["stability"] = (dg, (fnum(row.get("sigma_disp_eV")) or SCC.SIGMA_DISP_EV) * EV_KJMOL)
    elif str(row.get("disp_applicable")) == "True":
        miss.append("stability")
    lam = fnum(row.get("lambda_i_eV"))
    if lam is None or row.get("lambda_qc", "ok") != "ok": miss.append("kinetics")
    else: o["kinetics"] = (-lam, fnum(row.get("sigma_lambda_eV")) or SCC.SIGMA_LAMBDA_EV)
    return o, miss


def dominates(A, B):
    sh = set(A) & set(B)
    if len(sh) < 2:
        return False
    better = False
    for k in sh:
        (va, sa), (vb, sb) = A[k], B[k]
        tol = math.hypot(sa, sb)
        if vb - va > tol:
            return False
        if va - vb > tol:
            better = True
    return better


def fronts(rows):
    """{(scenario, pool, id): dict(pareto, partial, dominated_by, missing)}"""
    res = {}
    for scen, cap_key in SCC.CAPACITY_SCENARIOS.items():
        for pool in ("anolyte", "catholyte"):
            pc = [r for r in rows if r.get("pool") == pool]
            ob = {r["id"]: objectives(r, pool, cap_key) for r in pc}
            complete = [i for i, (_, ms) in ob.items() if not ms]
            for r in pc:
                i = r["id"]
                dom = sorted(j for j in complete if j != i and dominates(ob[j][0], ob[i][0]))
                res[(scen, pool, i)] = dict(pareto=(not ob[i][1] and not dom),
                                            partial=(bool(ob[i][1]) and not dom),
                                            dominated_by=";".join(dom), missing=";".join(ob[i][1]))
    return res


def audit_pareto(mine_sc):
    T = "pareto_shortlist"
    sc_pub = read_csv(RES / "scorecard.csv")
    cand_pub = [r for r in sc_pub if r.get("status") == "candidate" and str(r.get("rankable")) != "False"]
    f_pub = fronts(cand_pub)
    pub = {(r["scenario"], r["pool"], r["id"]): r for r in read_csv(RES / "pareto_shortlist.csv")}
    scmap = {(r["id"], r.get("pool")): r for r in sc_pub}
    for key in sorted(set(f_pub) | set(pub)):
        scen, pool, gid = key
        f, p = f_pub.get(key), pub.get(key)
        tag = f"{scen}/{pool}"
        if f is None or p is None:
            add(T, gid, f"{tag}:row_present", "present" if p else "", "present" if f else "", "logic")
            continue
        add(T, gid, f"{tag}:pareto_optimal", p["pareto_optimal"], str(f["pareto"]), "logic",
            note="front recomputed from PUBLISHED scorecard values")
        add(T, gid, f"{tag}:partial_front", p["partial_front"], str(f["partial"]), "logic")
        add(T, gid, f"{tag}:dominated_by", ";".join(sorted(filter(None, p["dominated_by"].split(";")))),
            f["dominated_by"], "logic")
        add(T, gid, f"{tag}:missing_objectives", p["missing_objectives"], f["missing"], "logic")
        s = scmap.get((gid, pool), {})
        cap_key = SCC.CAPACITY_SCENARIOS[scen]
        for pc, sc_col in (("E_V", "E_V"), ("capacity", cap_key), ("lambda_i_eV", "lambda_i_eV"),
                           ("dG_disp_kJmol", "dG_disp_kJmol"), ("n", "n_accessible")):
            add(T, gid, f"{tag}:{pc}_traces_scorecard", p.get(pc), fnum(s.get(sc_col)), "trace")
        sstat = s.get("status")
        if p["pareto_optimal"] == "True" and (sstat != "candidate" or f["missing"]):
            add(T, gid, f"{tag}:incomplete_on_front", "on front", "must not be", "logic")
    # INCOMPLETE scorecard rows must not appear on any primary front
    for r in sc_pub:
        if r.get("status") == "INCOMPLETE":
            on = [k for k, v in pub.items() if k[2] == r["id"] and v["pareto_optimal"] == "True"]
            add(T, r["id"], "INCOMPLETE_row_on_primary_front", "False", str(bool(on)), "logic")
    # front from OUR recomputed scorecard (propagation of upstream discrepancies)
    mine_rows = []
    for (gid, pool), m in mine_sc.items():
        if m.get("status") == "candidate" and m.get("rankable", True):
            mine_rows.append({**{k: ("" if v is None else str(v)) for k, v in m.items()},
                              "id": gid, "pool": pool})
    f_mine = fronts(mine_rows)
    for key in sorted(set(f_mine) | set(pub)):
        scen, pool, gid = key
        p, f = pub.get(key), f_mine.get(key)
        add(T, gid, f"{scen}/{pool}:pareto_optimal[from_recomputed_axes]",
            p["pareto_optimal"] if p else "", str(f["pareto"]) if f else "", "logic",
            note="front recomputed end-to-end from raw records")


# ============================================================================== sanity
def audit_multielectron():
    """Independently check consecutive reductions from the declared resting state."""
    pub = {r["id"]: r for r in read_csv(RES / "multielectron_stability.csv")}
    mine = {}
    factor = 8.314462618 * 298.15 / 96485.33212 * math.log(10)
    for gid, states in manifest_groups().items():
        rest = [r for r in states if str(r.get("n_e")) == "0"]
        if len(rest) != 1:
            continue
        q0 = int(rest[0]["charge"])
        rows = {r["q_ox"]: (event, r) for (g, event), r in MY_E.items()
                if g == gid and r.get("E") is not None}
        if q0 not in rows or q0 - 1 not in rows:
            continue
        ev1, r1 = rows[q0]; ev2, r2 = rows[q0 - 1]
        if r1["sR"] != r2["sO"]:
            continue
        E1, E2 = round(r1["E"], 3), round(r2["E"], 3)
        delta = round(E1 - E2, 3)
        lo, hi = ELE.WINDOW_V_VS_FC
        in1, in2 = lo <= E1 <= hi, lo <= E2 <= hi
        mine[gid] = dict(first_event=ev1, second_event=ev2, E1_vs_Fc_V=E1,
                         E2_vs_Fc_V=E2, dE12_V=delta, logK_comp=round(delta / factor, 1),
                         E1_in_window=str(in1), E2_in_window=str(in2),
                         usable_electrons=2 if in1 and in2 else (1 if in1 else 0))
    for gid in sorted(set(pub) | set(mine)):
        p, m = pub.get(gid), mine.get(gid)
        if p is None or m is None:
            add("multielectron_stability", gid, "row_present", "present" if p else "",
                "present" if m else "", "logic")
        else:
            for key, value in m.items():
                add("multielectron_stability", gid, key, p.get(key), value,
                    "logic" if isinstance(value, str) else "exact")


def audit_workbook():
    """Trace every exported table cell and detect spreadsheet error cells."""
    path = RES / "Candidates_computed.xlsx"
    if not path.exists():
        add("workbook", "all", "file_present", "", "present", "logic")
        return
    import openpyxl
    book = openpyxl.load_workbook(path, data_only=False, read_only=True)
    for name in ("scorecard", "pareto_shortlist", "redox_potentials", "reorganization"):
        if name not in book.sheetnames:
            add("workbook", name, "sheet_present", "", "present", "logic")
            continue
        cells = book[name].iter_rows(values_only=True)
        header = next(cells)
        data = list(cells); source = read_csv(RES / f"{name}.csv")
        add("workbook", name, "row_count", len(data), len(source), "trace")
        for index, (row, expected) in enumerate(zip(data, source), start=2):
            for column, cell in zip(header, row):
                raw = expected.get(column, "")
                numeric = fnum(raw)
                if numeric is not None and not isinstance(cell, bool):
                    add("workbook", name, f"{index}:{column}", cell, numeric, "trace", tol=1e-9)
                else:
                    add("workbook", name, f"{index}:{column}", "" if cell is None else str(cell), raw, "logic")
    errors = []
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.data_type == "e":
                    errors.append(f"{sheet.title}!{cell.coordinate}")
    add("workbook", "all", "error_cells", ";".join(errors), "", "logic")
    book.close()


SAN = []


def san(gid, st, check, value, threshold, basis, note=""):
    SAN.append(dict(id=gid, state=st, check=check, value=value, threshold=threshold,
                    threshold_basis=basis, note=note))


def vdw_volume_radius(xyz, spacing=0.2):
    """Radius (A) of the sphere with the same volume as the union of Bondi vdW spheres
    (RDKit GetRvdw), grid-integrated at `spacing` A."""
    from rdkit import Chem
    pt = Chem.GetPeriodicTable()
    sym, X = read_xyz(xyz)
    r = np.array([pt.GetRvdw(s) for s in sym])
    lo, hi = X.min(0) - r.max() - 0.5, X.max(0) + r.max() + 0.5
    axes = [np.arange(lo[i], hi[i], spacing) for i in range(3)]
    G = np.stack(np.meshgrid(*axes, indexing="ij"), -1).reshape(-1, 3)
    inside = np.zeros(len(G), bool)
    for xi, ri in zip(X, r):
        inside |= ((G - xi) ** 2).sum(1) <= ri * ri
    V = inside.sum() * spacing ** 3
    return (3 * V / (4 * math.pi)) ** (1 / 3)


def sanity():
    eps = ELE.SOLVENT["eps_r"]
    for gid in calc_ids():
        sts = calc_states(gid)
        neu = next((s for s in sts if s["charge"] == 0), None)
        for s in sts:
            st, q, m = s["state"], s["charge"], s["mult"]
            raw = s["raw_rec"]
            if raw is None:
                san(gid, st, "active_sp_record", "missing", "must exist", "contract",
                    "no ACTIVE-protocol energy record for current opt.xyz -> INCOMPLETE")
                continue
            for k in ("converged_smd", "converged_gas"):
                if raw.get(k) is not True:
                    san(gid, st, k, raw.get(k), "True", "contract", "unconverged SCF")
            if int(raw.get("charge", -99)) != q or int(raw.get("mult", -99)) != m:
                san(gid, st, "record_charge_mult", f"{raw.get('charge')}/{raw.get('mult')}", f"{q}/{m}", "contract")
            if s["g_th"] is None:
                san(gid, st, "g_thermal", "missing", "present", "contract")
            if (s["n_imag"] or 0) > 0:
                san(gid, st, "n_imag_thermal", s["n_imag"], 0, "contract",
                    f"xTB RRHO at DFT geometry; freq_min_cm={s['rj'].get('freq_min_cm')}")
            # spin contamination
            s2 = fnum(raw.get("gas_s_squared"))
            if s2 is not None:
                S = (m - 1) / 2
                exc = s2 - S * (S + 1)
                if abs(exc) > 0.1:
                    san(gid, st, "S2_excess_gas", round(exc, 4), 0.1, "heuristic",
                        f"<S^2>={s2:.4f}, ideal {S*(S+1):.3f} (mult {m})")
            # gas HOMO > 0 for an own-geometry anion: unbound in gas
            homo = fnum(raw.get("gas_homo_eV"))
            if q < 0 and homo is not None and homo > 0:
                san(gid, st, "gas_anion_unbound(HOMO>0)", round(homo, 3), 0.0, "computed",
                    "gas-phase anion at its own geometry has a positive HOMO (gas lambda/EA ill-defined)")
            # dG_solv plausibility
            dgs = fnum(raw.get("dG_solv_eV"))
            if dgs is None:
                continue
            if q == 0:
                if abs(dgs) > 1.0:
                    san(gid, st, "dG_solv_neutral", round(dgs, 3), "|x|>1.0 eV", "heuristic")
                continue
            a = vdw_volume_radius(s["dir"] / "opt.xyz")
            born = -(q * q) * COULOMB_EV_A / (2 * a) * (1 - 1 / eps)
            charging = dgs - (neu["dG_solv"] if neu and neu["dG_solv"] is not None else 0.0)
            ratio = charging / born
            basis_txt = ("dG_solv(ion) - dG_solv(neutral of same molecule)" if neu and neu["dG_solv"] is not None
                         else "dG_solv(ion) (no neutral partner)")
            if not (0.5 <= ratio <= 1.5):
                san(gid, st, "dG_solv_charging_vs_Born", round(charging, 3),
                    f"Born={born:.3f} eV (a={a:.2f} A); ratio {ratio:.2f} outside [0.5,1.5]",
                    "computed Born estimate; ratio window heuristic",
                    f"{basis_txt}; dG_solv={dgs:.3f} eV; a = radius of sphere with the Bondi "
                    f"vdW-union volume of opt.xyz; eps={eps}")
            if abs(q) == 1 and abs(dgs) > 3.5:
                san(gid, st, "dG_solv_monoion", round(dgs, 3), "|x|>3.5 eV", "heuristic")
        # gas IP / vertical IP vs -HOMO (Koopmans-type consistency for a range-separated hybrid)
        for O, R in adjacent(sts):
            if O["e_gas"] is None or R["e_gas"] is None:
                continue
            ip_ad = O["e_gas"] - R["e_gas"]
            cO = cross_path(gid, O["state"], R["state"], O["charge"], O["mult"])
            ip_v = None
            if cO.exists():
                d = json.loads(cO.read_text())
                if d.get("e_gas_eV") is not None:
                    ip_v = d["e_gas_eV"] - R["e_gas"]
            homo_R = fnum(R["rec"].get("gas_homo_eV"))
            if gid == "ferrocene":
                san(gid, f"{O['state']}/{R['state']}", "gas_IP_ferrocene",
                    f"adiabatic {ip_ad:.3f} eV; vertical {ip_v:.3f} eV" if ip_v is not None else f"adiabatic {ip_ad:.3f} eV",
                    "experimental gas IE: pending verification", "reported (no threshold)",
                    f"-HOMO(neutral) = {-homo_R:.3f} eV" if homo_R is not None else "")
            if ip_v is not None and homo_R is not None and abs(ip_v - (-homo_R)) > 1.0:
                san(gid, f"{O['state']}/{R['state']}", "gas_vertical_IP_vs_minus_HOMO",
                    round(ip_v, 3), f"-HOMO(R)={-homo_R:.3f} eV; |diff|>1.0 eV",
                    "heuristic (Koopmans-type, range-separated hybrid)",
                    f"vertical IP of R = E_O@geom_R - E_R@geom_R; diff {ip_v + homo_R:+.3f} eV")
            # cross-point spin contamination
            for c in (cO, cross_path(gid, R["state"], O["state"], R["charge"], R["mult"])):
                if c.exists():
                    d = json.loads(c.read_text())
                    s2 = fnum(d.get("gas_s_squared")); mm = int(d.get("mult", 1))
                    if s2 is not None and abs(s2 - ((mm - 1) / 2) * ((mm + 1) / 2)) > 0.1:
                        san(gid, c.stem.split("__")[0], "S2_excess_crosspoint",
                            round(s2 - ((mm - 1) / 2) * ((mm + 1) / 2), 4), 0.1, "heuristic")
                    if d.get("converged_gas") is not True:
                        san(gid, c.stem.split("__")[0], "crosspoint_converged_gas",
                            d.get("converged_gas"), "True", "contract")


# ================================================================================== main
def main():
    OUT.mkdir(parents=True, exist_ok=True)
    fc = audit_potentials()
    audit_orop()
    print(f"[fc] live Fc_abs = {fc:.6f} V (config record {ELE.FC_ABS_COMPUTED_V:.6f})")
    audit_disp()
    audit_multielectron()
    audit_reorg()
    audit_integrity()
    audit_capacity()
    mine_sc = audit_scorecard(fc)
    audit_pareto(mine_sc)
    audit_workbook()
    sanity()
    cols = ["table", "id", "key", "published", "recomputed", "abs_diff", "verdict",
            "check_type", "tol", "note"]
    with (OUT / "audit_recompute.csv").open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols); w.writeheader(); w.writerows(ROWS)
    with (OUT / "audit_sanity.csv").open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["id", "state", "check", "value", "threshold",
                                          "threshold_basis", "note"])
        w.writeheader(); w.writerows(SAN)
    # summary
    by = {}
    for r in ROWS:
        t = by.setdefault(r["table"], [0, 0])
        t[0] += 1; t[1] += r["verdict"] != "ok"
    print(f"{'table':32s} {'checked':>8s} {'discrep':>8s}")
    for t, (n, d) in by.items():
        print(f"{t:32s} {n:8d} {d:8d}")
    print(f"sanity flags: {len(SAN)}")
    print(f"wrote {OUT/'audit_recompute.csv'} and {OUT/'audit_sanity.csv'}")
    n_bad = sum(d for _, d in by.values())
    if n_bad:
        print(f"[FAIL] {n_bad} published value(s) do not match the independent recompute")
    return n_bad


if __name__ == "__main__":
    os.environ.setdefault("OMP_NUM_THREADS", "1")
    sys.exit(1 if main() else 0)
